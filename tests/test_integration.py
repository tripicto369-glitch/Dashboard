"""End-to-end test of the LHOS dashboard workbook in headless LibreOffice.

The expected results come from ``tests/oracle.py`` — an independent Python
model built from the user's requirements (it never reads the workbook's
formulas).  The scenario:

0. the shipped file itself (no LibreOffice): cached formula results (what
   Excel shows before it recalculates, e.g. in Protected View), button macros
   (``macro="[0]!…"`` on the drawing pictures), the month selector (no input
   tooltip, ▼ cell, selection in A1), names (``ui_Upd_<codename>`` on every
   tab, ``ui_Canvas``; no ``cfg_Company`` / ``ui_Updated`` / ``ui_Elapsed``),
   names used by the VBA code exist; then in LibreOffice: the cached results
   equal the computed empty state, the VBA modules equal ``src/vba``;
1. open a copy of ``dist/Дашборд_ЛХОС.xlsm`` -> empty state; threshold
   validation rules (green < yellow <= limit);
2. first import (mode 0) of the main mock file -> exact report, БД, sys_*
   cells, Расчет KPIs / sorted table / trends / max-per-day / month list and
   the visible texts of «Сводка» for TODAY (real date) and for several months
   chosen through ``sel_Month`` (incl. February with "н/д", March with the
   3.2 exceedance on 17.03 and the "0,45" text value, September with a blank
   last day); the rendered value column (PDF text) shows "*" for stale values;
   arrow / value colours (conditional formats evaluated in LibreOffice);
   the «Обновление: …» label and the freshness dot on every tab for several
   ages of the last import; other "today" dates (1st of a month, past month,
   exceedance day, "н/д" day) by replacing c_Today with a constant;
   EnsureMonthSelected prefers the current month;
3. threshold changes on «Настройки» re-zone everything;
4. append import (mode 2) of the update file -> counts, new object, the
   conflicting value is NOT overwritten; mode 0 on a non-empty base appends;
5. replace import (mode 1) of the main file -> base equals the main file again;
   a blank date row inside «БД» is tolerated by the formulas and by the next
   append; error paths (missing file, file without values, the dashboard
   itself, a period longer than the base holds) leave the base untouched;
   synthetic sources: December 2025 appended before the base, reordered
   columns, other spelling of a name, duplicate date rows, unmerged "Дата"
   header, text values, a small replace (no stray cells); 2-decimal rounding at
   the zone thresholds (1.504 green, 1.506 yellow …), near-ties in the sort;
   42 objects ("Показаны 24 из 42 СИКН", report warning, long file name);
   a name merged over both header rows, «Среднее»/«Итого» columns, the same
   dates on two sheets (reported), "<0,05"; a 9-year period (month list > 60
   months, months without values);
6. ClearDatabaseSilent (a Function returning True) -> empty state again;
7. the import log on «Настройки» has one row per import, newest first.

Cell addresses are discovered from the workbook (openpyxl), so the test does
not depend on the generated grid.  All mismatches are collected and printed;
the run fails if there is any.

LibreOffice-only workarounds (the workbook is correct in Excel; each is printed
as a "note", not a failure): Basic has no ``vbObjectError`` constant, so it is
defined in the working copy's modImport before the error paths are checked;
``Application.Calculation`` is per document in LibreOffice, so an import that
fails after the source workbook was closed leaves the dashboard in manual
calculation (and LibreOffice then stops evaluating conditional formats) — it is
switched back to automatic.  Conditional-format colours are checked by
evaluating the rules' formulas, not from the rendering.

Run:  ``python3 -I tests/test_integration.py``  (or via pytest).  Set
``LHOS_E2E_OUT=<dir>`` to keep the working copy and a JSON with mismatches,
``LHOS_XLSM=<path>`` to test another build of the workbook.
"""

from __future__ import annotations

import datetime as dt
import importlib.util
import json
import os
import re
import shutil
import sys
import tempfile
import time
import warnings
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tests"))

HAVE_LO = importlib.util.find_spec("uno") is not None and shutil.which("soffice") is not None
if not HAVE_LO and "pytest" in sys.modules:  # pragma: no cover
    import pytest

    pytest.skip("LibreOffice/uno not available", allow_module_level=True)

import openpyxl
import oracle as O
from openpyxl.formula.translate import Translator
from openpyxl.utils import column_index_from_string, get_column_letter

XLSM = Path(os.environ.get("LHOS_XLSM") or ROOT / "dist" / "Дашборд_ЛХОС.xlsm")
MAIN = ROOT / "dist" / "Мокап_ЛХОС_2026.xlsx"
UPD = ROOT / "dist" / "Мокап_ЛХОС_2026_дополнение.xlsx"
VBA_DIR = ROOT / "src" / "vba"

EPOCH = dt.date(1899, 12, 30)
EPOCH_DT = dt.datetime(1899, 12, 30)  # noqa: DTZ001 - Excel serials are local, naive
TOL = 1e-9

# Tabs with the common header (VBA code names of their sheets)
TAB_CODES = {"Сводка": "shDash", "СИКН": "shObj", "Тренды": "shTrend", "Тревоги": "shAlarm",
             "Отчеты": "shReport", "Настройки": "shSettings"}
SCRATCH_COL = "EZ"   # free column for evaluating helper formulas
TOO_LONG_TEXT = ("Слишком длинный период данных: 5845 дн. (допускается не более 3700).\n"
                 "Даты в данных: 01.01.2010 – 01.01.2026. Проверьте, нет ли опечатки в годе.")


def now_local() -> dt.datetime:
    """Local wall-clock time, as Excel's NOW() sees it (naive on purpose)."""
    return dt.datetime.now()  # noqa: DTZ005


def serial(d: dt.date) -> float:
    return float((d - EPOCH).days)


def from_serial(x) -> dt.date | None:
    if isinstance(x, (int, float)) and x > 0:
        return EPOCH + dt.timedelta(days=int(x))
    return None


def col_of(addr: str) -> str:
    return re.match(r"[A-Z]+", addr).group(0)


def row_of(addr: str) -> int:
    return int(re.search(r"\d+", addr).group(0))


# ---------------------------------------------------------------------------
# Layout discovery (addresses only; formulas are not evaluated or copied)
# ---------------------------------------------------------------------------

def _font_rgb(cell) -> str | None:
    c = cell.font.color if cell.font is not None else None
    return c.rgb if c is not None and isinstance(c.rgb, str) else None


def _cf_rules(ws) -> list:
    """[(sqref ranges, top-left, priority, formula, font rgb, num fmt)]."""
    out = []
    for cf in ws.conditional_formatting:
        ranges = list(cf.sqref.ranges)
        tl = ranges[0].coord.split(":")[0]
        for rule in cf.rules:
            rgb = num = None
            if rule.dxf is not None:
                if rule.dxf.font is not None and rule.dxf.font.color is not None:
                    rgb = rule.dxf.font.color.rgb
                if rule.dxf.numFmt is not None:
                    num = rule.dxf.numFmt.formatCode
            out.append((ranges, tl, rule.priority, rule.formula[0] if rule.formula else None, rgb, num))
    return out


def discover_layout(path: Path) -> dict:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(path, keep_vba=True)
    names = {k: v.attr_text for k, v in wb.defined_names.items()}

    def name_addr(n):
        sheet, addr = names[n].rsplit("!", 1)
        return sheet.strip("'"), addr.replace("$", "")

    ws = wb["Сводка"]
    cells = {}
    statics = {}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("="):
                cells.setdefault(c.value, []).append(c.coordinate)
            elif c.value is not None:
                statics.setdefault(str(c.value), []).append(c.coordinate)

    def find(pred, what):
        hits = [a for f, addrs in cells.items() if pred(f) for a in addrs]
        if not hits:
            raise AssertionError(f"layout: no cell for {what}")
        return hits

    def by_row(addrs):
        return sorted(addrs, key=lambda a: (row_of(a), column_index_from_string(col_of(a))))

    limits = find(lambda f: f == "=cfg_Limit", "limit")
    dash = {
        "today": find(lambda f: f == "=TODAY()", "today")[0],
        "data_date": find(lambda f: "c_DataDate" in f, "data date")[0],
        "loaded": find(lambda f: f.startswith("=IF(N(sys_LastImport)=0,\"—\""), "loaded")[0],
        "file": find(lambda f: f.startswith("=IF(sys_LastFile=\"\",\"—\""), "file")[0],
        "state_text": find(lambda f: f == "=c_StateText", "state text")[0],
        "state_sub": find(lambda f: f == "=c_StateSub", "state sub")[0],
        "max": find(lambda f: f.startswith('=IF(c_MaxVal="","—"'), "max")[0],
        "max_unit": find(lambda f: f.startswith('=IF(c_MaxVal="",""'), "max unit")[0],
        "max_name": find(lambda f: f == "=c_MaxName", "max name")[0],
        "avg": find(lambda f: f.startswith('=IF(c_AvgVal="","—"'), "avg")[0],
        "avg_unit": find(lambda f: f.startswith('=IF(c_AvgVal="",""'), "avg unit")[0],
        "avg_sub": find(lambda f: "c_CntValued" in f and "по " in f, "avg sub")[0],
        "green": find(lambda f: f == "=c_CntGreen", "green")[0],
        "yellow": find(lambda f: f == "=c_CntYellow", "yellow")[0],
        "red": find(lambda f: f == "=c_CntRed", "red")[0],
        "total": by_row(find(lambda f: f == "=c_ObjCount", "total"))[0],
        "donut_total": by_row(find(lambda f: f == "=c_ObjCount", "total"))[-1],
        "total_sub": find(lambda f: "c_CntNoData" in f, "no data")[0],
        "title": find(lambda f: "c_MonthLabel" in f, "title")[0],
        "note": find(lambda f: "c_StaleCount" in f, "note")[0],
        "footer": find(lambda f: "c_Footer" in f, "footer")[0],
        "source": find(lambda f: "Источник данных" in f, "source")[0],
        "leg_g": find(lambda f: f == "=c_LegG", "legend g")[0],
        "leg_y": find(lambda f: f == "=c_LegY", "legend y")[0],
        "leg_r": find(lambda f: f == "=c_LegR", "legend r")[0],
        "zone_r": find(lambda f: f == "=c_ZoneR", "zone r")[0],
        "zone_y": find(lambda f: f == "=c_ZoneY", "zone y")[0],
        "zone_g": find(lambda f: f == "=c_ZoneG", "zone g")[0],
        "limit": next(a for a in limits if "ppm" in ws[a].number_format),
        "scale_limit": next(a for a in limits if "ppm" not in ws[a].number_format),
        "scale_yellow": find(lambda f: f == "=cfg_Yellow", "scale yellow")[0],
        "scale_green": find(lambda f: f == "=cfg_Green", "scale green")[0],
        "sel": name_addr("sel_Month")[1],
    }

    r0 = 92  # first sorted row on «Расчет» (columns used by the dashboard rows)
    first = {
        "num": find(lambda f: f == f"=IF('Расчет'!$C${r0}=\"\",\"\",1)", "row num")[0],
        "name": find(lambda f: f == f"='Расчет'!$C${r0}", "row name")[0],
        "value": find(lambda f: f.startswith(f"=IF('Расчет'!$C${r0}=\"\",\"\",IF('Расчет'!$D${r0}"), "row value")[0],
        "status": find(lambda f: f"$E${r0}" in f and "НОРМА" in f, "row status")[0],
        "zone": find(lambda f: f"$E${r0}" in f and "c_ZoneG" in f, "row zone")[0],
        "change": find(lambda f: f"$F${r0}" in f and "↑" not in f, "row change")[0],
        "arrow": find(lambda f: f"$F${r0}" in f and "↑" in f, "row arrow")[0],
        "zone_h": find(lambda f: f == f"='Расчет'!$E${r0}", "zone helper")[0],
        "stale_h": find(lambda f: f == f"='Расчет'!$G${r0}", "stale helper")[0],
    }
    cols = {k: col_of(a) for k, a in first.items()}
    rows = []
    for k in range(O.TABLE_ROWS):
        a = find(lambda f, r=r0 + k: f == f"='Расчет'!$C${r}", f"table row {k}")[0]
        rows.append(row_of(a))

    # zone colours = font colours of the zone captions on the KPI card
    zone_rgb = {}
    for z, caption in ((1, "ЗЕЛЕНАЯ ЗОНА"), (2, "ЖЕЛТАЯ ЗОНА"), (3, "КРАСНАЯ ЗОНА")):
        zone_rgb[z] = _font_rgb(ws[statics[caption][0]])
    zone_rgb[4] = zone_rgb[3]

    # header of every tab: «Обновление: …» (ui_Upd_<code>) and the dot right of it
    upd, dots, words = {}, {}, {}
    for sheet, code in TAB_CODES.items():
        n = f"ui_Upd_{code}"
        if n not in names:
            raise AssertionError(f"layout: no name {n}")
        sh, addr = name_addr(n)
        assert sh == sheet, (n, sh)
        upd[sheet] = addr
        w = wb[sheet]
        r = row_of(addr)
        cand = [c for c in w[r] if c.value == "●" and c.column > column_index_from_string(col_of(addr))]
        dots[sheet] = cand[0].coordinate if cand else None
        words[sheet] = [c.coordinate for row in w.iter_rows(max_row=12) for c in row
                        if c.value == "TATNEFT"]

    st = wb["Настройки"]
    log_cols = [name_addr(f"log_C{j}")[1] for j in range(1, 10)]
    jcol = column_index_from_string(col_of(log_cols[0]))
    hdr_row = row_of(log_cols[0])
    hdr_last = max((m.max_row for m in st.merged_cells.ranges
                    if m.min_row == hdr_row and m.min_col == jcol), default=hdr_row)
    log_rows = sorted(m.min_row for m in st.merged_cells.ranges
                      if m.min_col == jcol and m.min_row > hdr_last)
    info = {}
    for row in st.iter_rows():
        for c in row:
            if not isinstance(c.value, str):
                continue
            if c.value.startswith("=COUNTA("):
                info["objects"] = c.coordinate
            elif c.value.startswith("=COUNT("):
                info["values"] = c.coordinate
            elif c.value.startswith("=IF(c_FirstDate=\"\""):
                info["period"] = c.coordinate

    calc = {n[2:]: name_addr(n)[1] for n in names if n.startswith("c_") and name_addr(n)[0] == "Расчет"}
    # month list: the rows counted by c_MonthCount (=COUNT(Bfirst:Blast))
    mc = wb["Расчет"][calc["MonthCount"]].value
    m = re.search(r"COUNT\(\$?B\$?(\d+):\$?B\$?(\d+)\)", mc)
    month_rows = (int(m.group(1)), int(m.group(2)))

    dv = {}
    for v in ws.data_validations.dataValidation:
        dv[str(v.sqref)] = v
    sdv = {}
    for v in st.data_validations.dataValidation:
        sdv[str(v.sqref)] = v

    return {
        "wb": wb, "names": names, "dash": dash, "cols": cols, "rows": rows,
        "cf": {sheet: _cf_rules(wb[sheet]) for sheet in TAB_CODES},
        "base_rgb": {k: _font_rgb(ws[f"{c}{rows[0]}"]) for k, c in cols.items()},
        "zone_rgb": zone_rgb, "upd": upd, "dots": dots, "words": words,
        "dot_rgb": {s: _font_rgb(wb[s][a]) if a else None for s, a in dots.items()},
        "log_cols": [col_of(a) for a in log_cols], "log_rows": log_rows, "info": info,
        "cfg": {n: name_addr(n)[1] for n in ("cfg_Green", "cfg_Yellow", "cfg_Limit")},
        "sys": {n: name_addr(n)[1] for n in ("sys_LastImport", "sys_LastFile", "sys_LastMode")},
        "calc": calc, "month_rows": month_rows, "dv": dv, "settings_dv": sdv,
    }


# ---------------------------------------------------------------------------
# Mismatch collector
# ---------------------------------------------------------------------------

class Checker:
    def __init__(self):
        self.fails: list[dict] = []
        self.notes: list[str] = []     # LibreOffice-only effects worked around (not failures)
        self.count = 0
        self.ctx = ""

    def eq(self, label, actual, expected, tol=TOL):
        self.count += 1
        ok = actual == expected
        if not ok and isinstance(actual, (int, float)) and isinstance(expected, (int, float)):
            ok = abs(actual - expected) <= tol
        if not ok:
            self.fails.append({"ctx": self.ctx, "what": label, "actual": actual, "expected": expected})
        return ok

    def true(self, label, cond, detail=""):
        self.count += 1
        if not cond:
            self.fails.append({"ctx": self.ctx, "what": label, "actual": detail, "expected": True})
        return cond


# ---------------------------------------------------------------------------
# The shipped file (no LibreOffice)
# ---------------------------------------------------------------------------

NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "rel": "http://schemas.openxmlformats.org/package/2006/relationships"}


def _sheet_parts(z: zipfile.ZipFile) -> dict:
    """{sheet name: part path} from workbook.xml and its relationships."""
    wbx = ET.fromstring(z.read("xl/workbook.xml"))
    rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
    target = {r.get("Id"): r.get("Target") for r in rels.findall("rel:Relationship", NS)}
    out = {}
    for s in wbx.find("m:sheets", NS):
        rid = s.get(f"{{{NS['r']}}}id")
        out[s.get("name")] = "xl/" + target[rid].lstrip("/").removeprefix("xl/")
    return out


def cached_results(path: Path) -> dict:
    """{sheet: {addr: (formula, cached value)}} for the tab sheets; cached value is a
    float, a str, ("ERR", text) or None (no <v>)."""
    out = {}
    with zipfile.ZipFile(path) as z:
        parts = _sheet_parts(z)
        for sheet in TAB_CODES:
            root = ET.fromstring(z.read(parts[sheet]))
            cells = {}
            for c in root.iter(f"{{{NS['m']}}}c"):
                f = c.find("m:f", NS)
                if f is None:
                    continue
                v = c.find("m:v", NS)
                t = c.get("t", "n")
                if v is None:
                    val = None
                elif t == "str":
                    val = v.text or ""
                elif t == "e":
                    val = ("ERR", v.text)
                elif t == "b":
                    val = v.text == "1"
                else:
                    val = float(v.text) if v.text else ""
                cells[c.get("r")] = (f.text or "", val)
            out[sheet] = cells
    return out


def drawing_macros(path: Path) -> dict:
    """{sheet: {shape name: macro attribute}} for the button pictures."""
    out = {}
    with zipfile.ZipFile(path) as z:
        parts = _sheet_parts(z)
        for sheet, part in parts.items():
            rel_path = part.replace("worksheets/", "worksheets/_rels/") + ".rels"
            if rel_path not in z.namelist():
                continue
            rels = ET.fromstring(z.read(rel_path))
            for r in rels.findall("rel:Relationship", NS):
                if not r.get("Target", "").endswith(".xml") or "drawing" not in r.get("Target"):
                    continue
                dpath = "xl/drawings/" + r.get("Target").rsplit("/", 1)[1]
                xml = z.read(dpath).decode("utf-8")
                for m in re.finditer(r'<xdr:(pic|sp)( macro="([^"]*)")?[^>]*>\s*<xdr:nv(?:Pic|Sp)Pr>\s*'
                                     r'<xdr:cNvPr [^>]*name="(btn\w+)"', xml):
                    out.setdefault(sheet, {})[m.group(4)] = m.group(3)
    return out


def check_static(path: Path, L: dict, ck: Checker):
    ck.ctx = "shipped file"
    names = L["names"]
    for gone in ("ui_Updated", "ui_Elapsed", "cfg_Company"):
        ck.true(f"no name {gone}", gone not in names, names.get(gone))
    for n in ("ui_Canvas", "ui_MonthArrow", "sel_Month", "lst_Months", "m_Labels"):
        ck.true(f"name {n} exists", n in names)
    # every tab: «Обновление» name, freshness dot with 3 rules on sys_LastImport, wordmark
    for sheet in TAB_CODES:
        ck.true(f"{sheet}: freshness dot right of ui_Upd", L["dots"][sheet] is not None)
        rules = [r for r in L["cf"][sheet] if L["dots"][sheet] in {c for rg in r[0] for c in _cells(rg.coord)}]
        ck.eq(f"{sheet}: dot rules", len(rules), 3)
        ck.true(f"{sheet}: dot rules use sys_LastImport",
                all("sys_LastImport" in (r[3] or "") for r in rules), str([r[3] for r in rules]))
        ck.eq(f"{sheet}: wordmark TATNEFT in the header", len(L["words"][sheet]), 1)
    # canvas covers the whole dashboard drawing (A1 .. beyond the footer)
    sh, rng = names["ui_Canvas"].rsplit("!", 1)
    a, b = rng.replace("$", "").split(":")
    ck.eq("ui_Canvas sheet", sh.strip("'"), "Сводка")
    ck.eq("ui_Canvas starts at A1", a, "A1")
    for key in ("footer", "today", "source", "donut_total"):
        addr = L["dash"][key]
        ck.true(f"ui_Canvas covers {key} ({addr})",
                row_of(addr) <= row_of(b) and column_index_from_string(col_of(addr))
                <= column_index_from_string(col_of(b)), b)
    # month selector: list validation, no input tooltip, ▼ right next to it
    ws = L["wb"]["Сводка"]
    sel = L["dash"]["sel"]
    sel_m = [m for m in ws.merged_cells.ranges if m.min_row == row_of(sel)
             and get_column_letter(m.min_col) == col_of(sel)]
    sel_dv = [v for k, v in L["dv"].items() if sel in _cells(k.split()[0])]
    ck.eq("sel_Month has one validation", len(sel_dv), 1)
    if sel_dv:
        v = sel_dv[0]
        ck.eq("sel_Month validation type", v.type, "list")
        ck.eq("sel_Month list source", v.formula1, "lst_Months")
        ck.true("sel_Month: no input message", not v.prompt and not v.promptTitle,
                f"{v.promptTitle!r} {v.prompt!r}")
        ck.true("sel_Month: stop on invalid input", v.showErrorMessage and v.errorStyle in (None, "stop"),
                f"{v.showErrorMessage} {v.errorStyle}")
    arrow = names["ui_MonthArrow"].rsplit("!", 1)[1].replace("$", "")
    if sel_m:
        ck.eq("ui_MonthArrow right of sel_Month (row)", row_of(arrow), sel_m[0].min_row)
        ck.eq("ui_MonthArrow right of sel_Month (column)", column_index_from_string(col_of(arrow)),
              sel_m[0].max_col + 1)
    ck.eq("ui_MonthArrow shows ▼", ws[arrow].value, "▼")
    selection = ws.sheet_view.selection[0].activeCell if ws.sheet_view.selection else None
    ck.true("Сводка opens with A1 selected (no tooltip over the KPIs)", selection in (None, "A1"),
            selection)
    # threshold inputs: custom ordering rules
    for n, addr in L["cfg"].items():
        hits = [v for k, v in L["settings_dv"].items() if addr in _cells(k.split()[0])]
        ck.true(f"{n}: custom validation rule", len(hits) == 1 and hits[0].type == "custom",
                [h.type for h in hits])
    # buttons carry their macro in the file (no VBA needed to assign it)
    macros = drawing_macros(path)
    ck.eq("buttons on «Сводка»", macros.get("Сводка"),
          {"btnImportData": "[0]!ImportData", "btnExitApp": "[0]!ExitApp"})
    ck.eq("buttons on «Настройки»", macros.get("Настройки"),
          {"btnImportData": "[0]!ImportData", "btnClearDatabase": "[0]!ClearDatabase"})
    # names used by the VBA code exist in the workbook
    for f in sorted(VBA_DIR.glob("*.*")):
        for lit in re.findall(r'Names\("([^"]+)"\)', f.read_text(encoding="utf-8")):
            ck.true(f"{f.name}: Names(\"{lit}\") defined", lit in names)
    # cached results: what Excel shows before recalculation (Protected View)
    cache = cached_results(path)
    dash = cache["Сводка"]
    for key, want in (("state_text", "НЕТ ДАННЫХ"), ("state_sub", "Загрузите файл с данными"),
                      ("note", "База пуста — нажмите «ИМПОРТ ДАННЫХ»"), ("max", "—"),
                      ("avg", "—"), ("data_date", "—"), ("loaded", "—"), ("file", "—"),
                      ("source", "Данные не загружены — нажмите «Импорт данных»"),
                      ("footer", "СОСТОЯНИЕ: НЕТ ДАННЫХ"), ("max_unit", ""), ("max_name", "")):
        ck.eq(f"cached {key}", dash.get(L["dash"][key], (None, None))[1], want)
    ck.eq("cached total", dash.get(L["dash"]["total"], (None, None))[1], 0.0)
    for sheet in TAB_CODES:
        ck.eq(f"cached «Обновление» on {sheet}", cache[sheet].get(L["upd"][sheet], (None, None))[1],
              "Данные еще не загружались")
        todays = [v for f, v in cache[sheet].values() if f == "TODAY()"]
        ck.true(f"cached TODAY() on {sheet} is a date", len(todays) == 1 and isinstance(todays[0], float)
                and todays[0] > 45000, todays)
    title = dash.get(L["dash"]["title"], (None, None))[1]
    ck.true("cached title", isinstance(title, str) and title.startswith("СОДЕРЖАНИЕ ЛХОС ПО СИКН — "),
            title)
    for k in range(O.TABLE_ROWS):
        r = L["rows"][k]
        got = tuple(dash.get(f"{L['cols'][c]}{r}", (None, None))[1]
                    for c in ("num", "name", "value", "change", "arrow"))
        ck.true(f"cached table row {k + 1} empty", all(x == "" for x in got), got)
    return cache


def _cells(rng: str) -> set:
    """Cell addresses of an A1 range string (single range)."""
    from openpyxl.utils.cell import range_boundaries
    c1, r1, c2, r2 = range_boundaries(str(rng))
    return {f"{get_column_letter(c)}{r}" for c in range(c1, c2 + 1) for r in range(r1, r2 + 1)}


VOLATILE = ("TODAY()", "NOW()", "c_MonthLabel")


def check_cache_matches_lo(doc, L, cache: dict, ck: Checker):
    """Every cached formula result on the tabs equals LibreOffice's empty-state result
    (except today-dependent cells and the helper cells right of the «Сводка» canvas,
    which only feed conditional formats)."""
    ck.ctx = "cached results vs LibreOffice (empty book)"
    canvas = L["names"]["ui_Canvas"].rsplit("!", 1)[1].replace("$", "").split(":")[1]
    last_col = column_index_from_string(col_of(canvas))
    bad = []
    n = 0
    for sheet, cells in cache.items():
        for addr, (f, val) in cells.items():
            if any(x in f for x in VOLATILE):
                continue
            if sheet == "Сводка" and column_index_from_string(col_of(addr)) > last_col:
                continue
            c = doc.cell(sheet, addr)
            if c.getError():
                lo = ("ERR",)
            elif c.FormulaResultType2 == 2:
                lo = c.getString()
            else:
                lo = c.getValue()
            n += 1
            same = (val == lo or (isinstance(val, float) and isinstance(lo, float) and abs(val - lo) < 1e-9)
                    or (isinstance(val, tuple) and isinstance(lo, tuple)))
            if not same:
                bad.append((sheet, addr, f[:60], val, lo))
    ck.true(f"cached results equal LibreOffice ({n} cells)", not bad, str(bad[:10]))
    ck.true("cached results checked", n > 100, n)


def check_vba_sources(doc, ck: Checker):
    """The VBA project in the workbook is the code in src/vba (dist is not stale)."""
    ck.ctx = "VBA project vs src/vba"

    def norm(s: str) -> list:
        lines = [ln.rstrip() for ln in s.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
        lines = [ln for ln in lines if not ln.startswith(("Rem Attribute", "Attribute ", "Option VBASupport"))]
        while lines and not lines[-1]:
            lines.pop()
        while lines and not lines[0]:
            lines.pop(0)
        return lines

    mods = {}
    for lib in doc.basic_modules().values():
        mods.update(lib)
    for f in sorted(VBA_DIR.glob("*.*")):
        name = f.stem
        ck.true(f"module {name} in the workbook", name in mods)
        if name in mods:
            a, b = norm(mods[name]), norm(f.read_text(encoding="utf-8"))
            diff = next((i for i, (x, y) in enumerate(zip(a, b)) if x != y), None)
            if diff is None and len(a) != len(b):
                diff = min(len(a), len(b))
            ck.true(f"module {name} equals src/vba/{f.name}", diff is None,
                    None if diff is None else f"line {diff + 1}: {a[diff:diff + 1]} vs {b[diff:diff + 1]}")


# ---------------------------------------------------------------------------
# Reading the workbook through UNO
# ---------------------------------------------------------------------------

def cell_value(doc, sheet, addr):
    """None (empty / ""), float, str or ("ERR", code)."""
    c = doc.cell(sheet, addr)
    if c.getError():
        return ("ERR", c.getError())
    v = doc.get(sheet, addr)
    if v == "":
        return None
    return v


def read_block(doc, sheet, rng):
    """2-D list of cell_value for a range (errors as ("ERR", code))."""
    s = doc.sheet(sheet)
    r = s.getCellRangeByName(rng)
    data = r.getDataArray()
    out = []
    for i, row in enumerate(data):
        orow = []
        for j, v in enumerate(row):
            if v is None or v == "":  # getDataArray gives None for error results
                cell = r.getCellByPosition(j, i)
                err = cell.getError()
                orow.append(("ERR", err) if err else None)
            else:
                orow.append(v)
        out.append(orow)
    return out


def calc_param(doc, L, name):
    return cell_value(doc, "Расчет", L["calc"][name])


def api_formula(f: str) -> str:
    """Excel formula text -> LibreOffice API grammar (";" between arguments)."""
    out, q = [], False
    for ch in f:
        if ch == '"':
            q = not q
        out.append(";" if ch == "," and not q else ch)
    s = "".join(out)
    return s if s.startswith("=") else "=" + s


def eval_formulas(doc, sheet: str, formulas: list) -> list:
    """Evaluate formulas on a sheet (in a free column); TRUE/FALSE -> bool, error -> None."""
    s = doc.sheet(sheet)
    cells = []
    for i, f in enumerate(formulas):
        c = s.getCellRangeByName(f"{SCRATCH_COL}{i + 1}")
        c.setFormula(api_formula(f))
        cells.append(c)
    doc.recalc()
    out = [None if c.getError() else bool(c.getValue()) for c in cells]
    if formulas:
        s.getCellRangeByName(f"{SCRATCH_COL}1:{SCRATCH_COL}{len(formulas)}").clearContents(1 | 2 | 4 | 16)
        doc.recalc()
    return out


def effective_colors(doc, L, sheet: str, addrs: list, base: dict | None = None) -> list:
    """Font colour of cells after conditional formatting (highest-priority true rule
    with a font colour, else the cell's own colour)."""
    ws = L["wb"][sheet]
    jobs = []
    for addr in addrs:
        for ranges, tl, prio, f, rgb, _num in L["cf"][sheet]:
            if f is None or rgb is None:
                continue
            if any(addr in _cells(r.coord) for r in ranges):
                jobs.append((addr, prio, rgb, Translator("=" + f, origin=tl).translate_formula(addr)))
    res = eval_formulas(doc, sheet, [j[3] for j in jobs])
    out = []
    for addr in addrs:
        hits = sorted((prio, rgb) for (a, prio, rgb, _f), ok in zip(jobs, res) if a == addr and ok)
        own = (base or {}).get(addr) or _font_rgb(ws[addr])
        out.append(hits[0][1] if hits else own)
    return out


# ---------------------------------------------------------------------------
# Comparisons
# ---------------------------------------------------------------------------

def check_view(doc, L, v: O.View, ck: Checker, real_today: bool = True):
    """Compare «Расчет» and the visible texts of «Сводка» with the oracle view."""
    P = lambda n: calc_param(doc, L, n)
    ck.eq("c_MonthStart", from_serial(P("MonthStart")), v.month_start)
    ck.eq("c_MonthEnd", from_serial(P("MonthEnd")), v.month_end)
    ck.eq("c_RefDate", from_serial(P("RefDate")), v.ref_date)
    ck.eq("c_IsCurrent", bool(P("IsCurrent")), v.is_current)
    ck.eq("c_MonthLabel", P("MonthLabel"), v.label)
    ck.eq("c_ObjCount", P("ObjCount"), float(v.obj_count))
    ck.eq("c_DataDate", from_serial(P("DataDate")), v.data_date)
    ck.eq("c_CntValued", P("CntValued"), float(v.cnt_valued))
    ck.eq("c_CntGreen", P("CntGreen"), float(v.cnt_green))
    ck.eq("c_CntYellow", P("CntYellow"), float(v.cnt_yellow))
    ck.eq("c_CntRed", P("CntRed"), float(v.cnt_red))
    ck.eq("c_CntOver", P("CntOver"), float(v.cnt_over))
    ck.eq("c_CntNoData", P("CntNoData"), float(v.cnt_nodata))
    ck.eq("c_MaxVal", P("MaxVal"), v.max_val)
    ck.eq("c_MaxName", P("MaxName"), v.max_name or None)
    ck.eq("c_AvgVal", P("AvgVal"), v.avg_val, tol=1e-9)
    ck.eq("c_State", P("State"), float(v.state))
    ck.eq("c_StaleCount", P("StaleCount"), float(v.stale_count))
    ck.eq("c_LastMaxDate", from_serial(P("LastMaxDate")), v.last_max_date)

    # objects in base order (rows 46..85)
    blk = read_block(doc, "Расчет", "A46:K85")
    for i in range(O.MAX_OBJ):
        row = blk[i]
        if i < len(v.objects):
            o = v.objects[i]
            tag = f"obj[{i + 1}] {o.name}"
            ck.eq(tag + " name", row[1], o.name)
            ck.eq(tag + " group", row[2], o.group or None)
            ck.eq(tag + " value date", from_serial(row[3]), o.value_date)
            ck.eq(tag + " value", row[4], o.value)
            ck.eq(tag + " prev date", from_serial(row[5]), o.prev_date)
            ck.eq(tag + " prev value", row[6], o.prev_value)
            ck.eq(tag + " change", row[7], o.change, tol=1e-9)
            ck.eq(tag + " zone", row[8], float(o.zone))
            ck.eq(tag + " stale", row[10], 1.0 if o.stale else 0.0)
        else:
            ck.eq(f"obj[{i + 1}] empty name", row[1], None)

    # sorted table + trends (rows 92..131, trend L..AP)
    blk = read_block(doc, "Расчет", "A92:AP131")
    for k in range(O.MAX_OBJ):
        row = blk[k]
        if k < len(v.sorted):
            o = v.sorted[k]
            tag = f"sorted[{k + 1}]"
            ck.eq(tag + " name", row[2], o.name)
            ck.eq(tag + " value", row[3], o.value)
            ck.eq(tag + " zone", row[4], float(o.zone))
            ck.eq(tag + " change", row[5], o.change, tol=1e-9)
            ck.eq(tag + " stale", row[6], 1.0 if o.stale else 0.0)
            trend = [None if isinstance(x, tuple) else x for x in row[11:42]]
            errs_ok = all(isinstance(x, tuple) for x, e in zip(row[11:42], v.trend[k]) if e is None)
            ck.eq(tag + f" trend {o.name}", trend, v.trend[k])
            ck.true(tag + " trend gaps are #N/A", errs_ok, str(row[11:42]))
        else:
            ck.eq(f"sorted[{k + 1}] empty name", row[2], None)
            ck.true(f"sorted[{k + 1}] trend all #N/A",
                    all(isinstance(x, tuple) for x in row[11:42]), str(row[11:42])[:120])

    # day dates row 91 and max-per-day row 134, last point row 138
    days = read_block(doc, "Расчет", "L91:AP91")[0]
    exp_days = [v.days[k] if k < len(v.days) else None for k in range(31)]
    ck.eq("day dates", [from_serial(x) for x in days], exp_days)
    mx = read_block(doc, "Расчет", "L134:AP134")[0]
    ck.eq("max per day", [None if isinstance(x, tuple) else x for x in mx], v.maxday)
    last = read_block(doc, "Расчет", "L138:AP138")[0]
    exp_last = [v.maxday[k] if (v.last_max_date and k < len(v.days) and v.days[k] == v.last_max_date)
                else None for k in range(31)]
    ck.eq("last point", [None if isinstance(x, tuple) else x for x in last], exp_last)
    lim = read_block(doc, "Расчет", "L135:AP137")
    for r, val, nm in ((0, v.cfg.limit, "limit"), (1, v.cfg.yellow, "yellow"), (2, v.cfg.green, "green")):
        ck.eq(f"threshold line {nm}", [None if isinstance(x, tuple) else x for x in lim[r]],
              [val if k < len(v.days) else None for k in range(31)])

    # month list (all rows counted by c_MonthCount)
    r1, r2 = L["month_rows"]
    ml = read_block(doc, "Расчет", f"B{r1}:C{r2}")
    got = [(from_serial(a), b) for a, b in ml if b]
    ck.eq("month list", got, v.months)
    ck.eq("c_MonthCount", P("MonthCount"), float(len(v.months)))

    # visible texts on «Сводка»
    exp = O.dash_texts(v)
    for key, want in exp.items():
        addr = L["dash"][key]
        ck.eq(f"Сводка {key} ({addr})", doc.get_string("Сводка", addr), want)
    if real_today:  # the header always shows the real date (TODAY())
        ck.eq("Сводка today", doc.get_string("Сводка", L["dash"]["today"]), v.today.strftime("%d.%m.%Y"))
    texts = O.table_texts(v)
    c = L["cols"]
    for k, r in enumerate(L["rows"]):
        got = tuple(doc.get_string("Сводка", f"{c[n]}{r}")
                    for n in ("num", "name", "value", "status", "zone", "change", "arrow"))
        ck.eq(f"Сводка table row {k + 1} (row {r})", got, texts[k])
        o = v.sorted[k] if k < len(v.sorted) else None
        zh = cell_value(doc, "Сводка", f"{c['zone_h']}{r}")
        sh = cell_value(doc, "Сводка", f"{c['stale_h']}{r}")
        if o is not None:
            ck.eq(f"Сводка row {k + 1} zone helper", zh, float(o.zone))
            ck.eq(f"Сводка row {k + 1} stale helper", sh, 1.0 if o.stale else 0.0)


def check_colors(doc, L, v: O.View, ck: Checker):
    """Value and arrow colours of the table rows (conditional formats evaluated):
    value — colour of its zone; arrow ↑/↓ — colour of the zone, «—» — neutral."""
    zone_rgb = L["zone_rgb"]
    zone_set = set(zone_rgb.values())
    c = L["cols"]
    n = min(len(v.sorted), O.TABLE_ROWS)
    rows = L["rows"][:n]
    vals = effective_colors(doc, L, "Сводка", [f"{c['value']}{r}" for r in rows])
    arrows = effective_colors(doc, L, "Сводка", [f"{c['arrow']}{r}" for r in rows])
    for k in range(n):
        o = v.sorted[k]
        if o.zone:
            ck.eq(f"value colour row {k + 1} {o.name} (zone {o.zone})", vals[k], zone_rgb[o.zone])
        else:
            ck.true(f"value colour row {k + 1} {o.name} (no data) is not a zone colour",
                    vals[k] not in zone_set, vals[k])
        a = O.arrow_of(o.change)
        if a in ("↑", "↓"):
            ck.eq(f"arrow {a} colour row {k + 1} {o.name}", arrows[k], zone_rgb[o.zone])
        elif a == "—":
            ck.true(f"arrow — colour row {k + 1} {o.name} is neutral", arrows[k] not in zone_set,
                    arrows[k])


def read_db(doc):
    """(objects, groups, {(date, name): value}, dates list) from «БД»."""
    s = doc.sheet("БД")
    hdr = s.getCellRangeByName("B1:CV2").getDataArray()
    names = []
    for x in hdr[0]:
        if x == "":
            break
        names.append(x)
    groups = {n: hdr[1][i] for i, n in enumerate(names)}
    col = s.getCellRangeByName("A3:A3702").getDataArray()
    dates = []
    for (x,) in col:
        if x == "":
            break
        dates.append(from_serial(x))
    vals = {}
    if names and dates:
        last = get_column_letter(1 + len(names))
        data = s.getCellRangeByName(f"B3:{last}{2 + len(dates)}").getDataArray()
        for i, row in enumerate(data):
            for j, x in enumerate(row):
                if x != "":
                    vals[(dates[i], names[j])] = x
    # anything outside the block (rows below, header/values right of the objects)?
    extra = s.getCellRangeByName(f"A{3 + len(dates)}:CV{3 + len(dates) + 5}").getDataArray()
    stray = any(x != "" for row in extra for x in row)
    right = s.getCellRangeByName(f"{get_column_letter(2 + len(names))}1:CV{max(3, 2 + len(dates))}").getDataArray()
    stray = stray or any(x != "" for row in right for x in row)
    return names, groups, vals, dates, stray


def check_db(doc, db: O.Db, ck: Checker):
    names, groups, vals, dates, stray = read_db(doc)
    ck.eq("БД objects", names, db.objects)
    ck.eq("БД groups", groups, {n: db.groups.get(n, "") for n in db.objects})
    ck.eq("БД dates contiguous first..last", dates, db.dates())
    ck.eq("БД value count", len(vals), len(db.values))
    diff = {k: (vals.get(k), db.values.get(k)) for k in set(vals) | set(db.values)
            if vals.get(k) is None or db.values.get(k) is None or abs(vals[k] - db.values[k]) > 1e-12}
    ck.true("БД values equal oracle", not diff,
            str(sorted(((d.isoformat(), n), a, e) for (d, n), (a, e) in diff.items())[:10]))
    ck.true("БД no stray cells below / right of the data", not stray)


def norm_report(text: str) -> str:
    s = str(text).replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"(?<=\d)[\u00a0\u202f ](?=\d{3}(?!\d))", " ", s)


def check_report(res, rep: O.ImportReport, src: O.Source, ck: Checker):
    res = str(res)
    ck.true("import returns OK", res.startswith("OK\n"), res[:200])
    got = norm_report(res.partition("\n")[2])
    want = O.report_text(rep, src)
    if got != want:
        gl, wl = got.split("\n"), want.split("\n")
        i = next((i for i, (a, b) in enumerate(zip(gl, wl)) if a != b), min(len(gl), len(wl)))
        ck.eq(f"report text (first difference at line {i + 1})", gl[i:i + 3], wl[i:i + 3])
    else:
        ck.eq("report text", got, want)
    # summary first, details after a blank line
    lines = got.split("\n")
    if "" in lines:
        blank = lines.index("")
        ck.true("report: period in the summary part",
                any(ln.startswith("Период данных в базе:") for ln in lines[:blank]))
        ck.true("report: file statistics after the blank line",
                lines[blank + 1].startswith("Значений в файле:") if blank + 1 < len(lines) else False)
    else:
        ck.true("report has a blank line between summary and details", False, got[:300])


def check_sys(doc, L, ck: Checker, *, file_name, mode_text, since: dt.datetime | None):
    s = L["sys"]
    last = cell_value(doc, "Настройки", s["sys_LastImport"])
    if since is None:
        ck.eq("sys_LastImport empty", last, None)
    else:
        ok = isinstance(last, float)
        if ok:
            ts = EPOCH_DT + dt.timedelta(days=last)
            ok = since - dt.timedelta(minutes=2) <= ts <= now_local() + dt.timedelta(minutes=2)
        ck.true("sys_LastImport ~ now", ok, str(last))
    ck.eq("sys_LastFile", cell_value(doc, "Настройки", s["sys_LastFile"]), file_name)
    ck.eq("sys_LastMode", cell_value(doc, "Настройки", s["sys_LastMode"]), mode_text)
    d = L["dash"]
    for sheet, addr in L["upd"].items():
        txt = doc.get_string(sheet, addr)
        if since is None:
            ck.eq(f"{sheet} «Обновление» text", txt, "Данные еще не загружались")
        else:
            ck.true(f"{sheet} «Обновление» just now",
                    txt in ("Обновление: только что", "Обновление: 1 мин назад"), txt)
    if since is None:
        ck.eq("Сводка loaded", doc.get_string("Сводка", d["loaded"]), "—")
    else:
        ck.true("Сводка loaded dd.mm.yyyy hh:mm",
                re.fullmatch(r"\d\d\.\d\d\.\d{4} \d\d:\d\d", doc.get_string("Сводка", d["loaded"])) is not None,
                doc.get_string("Сводка", d["loaded"]))
    ck.eq("Сводка file", doc.get_string("Сводка", d["file"]), O.file_label(file_name))
    ck.eq("Сводка source", doc.get_string("Сводка", d["source"]), O.source_label(file_name))


def check_info(doc, L, db: O.Db, ck: Checker):
    i = L["info"]
    ck.eq("Настройки objects in base", cell_value(doc, "Настройки", i["objects"]), float(len(db.objects)))
    ck.eq("Настройки values in base", cell_value(doc, "Настройки", i["values"]), float(len(db.values)))
    want = "—" if db.empty else f"{O.fmt_date(db.first)} – {O.fmt_date(db.last)}"
    ck.eq("Настройки period", doc.get_string("Настройки", i["period"]), want)


def read_log(doc, L):
    rows = []
    for r in L["log_rows"]:
        vals = [cell_value(doc, "Настройки", f"{c}{r}") for c in L["log_cols"]]
        if all(v is None for v in vals):
            continue
        rows.append(vals)
    return rows


def set_month(doc, L, label):
    doc.cell("Сводка", L["dash"]["sel"]).setString(label)
    doc.recalc()


def set_cfg(doc, L, cfg: O.Cfg):
    for n, v in (("cfg_Green", cfg.green), ("cfg_Yellow", cfg.yellow), ("cfg_Limit", cfg.limit)):
        doc.cell("Настройки", L["cfg"][n]).setValue(v)
    doc.recalc()


def check_freshness(doc, L, ck: Checker, delta: dt.timedelta | None):
    """«Обновление: …» on every tab and the colour of the freshness dot."""
    want_txt = O.elapsed_text(delta)
    want_col = O.freshness(delta)
    names = {"green": L["zone_rgb"][1], "yellow": L["zone_rgb"][2], "red": L["zone_rgb"][3]}
    for sheet in TAB_CODES:
        ck.eq(f"{sheet} «Обновление» for {delta}", doc.get_string(sheet, L["upd"][sheet]), want_txt)
        dot = L["dots"][sheet]
        if dot is None:
            continue
        col = effective_colors(doc, L, sheet, [dot])[0]
        if want_col is None:
            ck.true(f"{sheet} dot colour (never imported) is not a zone colour",
                    col not in set(names.values()), col)
        else:
            ck.eq(f"{sheet} dot colour for {delta}", col, names[want_col])


def check_elapsed(doc, L, ck: Checker):
    addr = L["sys"]["sys_LastImport"]
    keep = doc.cell("Настройки", addr).getValue()
    for delta in (dt.timedelta(seconds=20), dt.timedelta(minutes=5, seconds=20),
                  dt.timedelta(hours=2, minutes=7, seconds=20),
                  dt.timedelta(hours=23, minutes=58, seconds=30),
                  dt.timedelta(days=1, hours=2, minutes=1, seconds=20),
                  dt.timedelta(days=2, hours=23, minutes=58, seconds=30),
                  dt.timedelta(days=3, hours=4, minutes=10, seconds=20)):
        stamp = now_local() - delta
        doc.cell("Настройки", addr).setValue((stamp - EPOCH_DT).total_seconds() / 86400)
        doc.recalc()
        check_freshness(doc, L, ck, delta)
    doc.cell("Настройки", addr).setValue(keep)
    doc.recalc()


def check_threshold_rules(doc, L, ck: Checker):
    """Validation of the threshold inputs: 0 <= green < yellow <= limit."""
    rules = {}
    for n, addr in L["cfg"].items():
        hits = [v for k, v in L["settings_dv"].items() if addr in _cells(k.split()[0])]
        if hits and hits[0].formula1:
            rules[n] = hits[0].formula1
    if len(rules) != 3:
        ck.true("threshold rules found", False, list(rules))
        return
    cases = [(1.5, 2.3, 3.0), (2.5, 2.3, 3.0), (1.5, 2.3, 2.3), (1.5, 2.3, 2.0), (-0.1, 2.3, 3.0),
             (1.5, 1.5, 3.0), (0.0, 0.1, 0.1), ("abc", 2.3, 3.0), (1.25, 2.35, 3.05)]
    for g, y, lim in cases:
        for n, v in (("cfg_Green", g), ("cfg_Yellow", y), ("cfg_Limit", lim)):
            c = doc.cell("Настройки", L["cfg"][n])
            c.setString(v) if isinstance(v, str) else c.setValue(v)
        res = dict(zip(rules, eval_formulas(doc, "Настройки", list(rules.values()))))
        num = lambda x: not isinstance(x, str)
        want = {"cfg_Green": num(g) and num(y) and g >= 0 and g < y,
                "cfg_Yellow": num(y) and num(g) and num(lim) and g < y <= lim,
                "cfg_Limit": num(lim) and num(y) and lim >= y}
        for n in rules:
            ck.eq(f"{n} accepted for green={g}, yellow={y}, limit={lim}", res[n], want[n])
    set_cfg(doc, L, O.Cfg())


def check_pdf_values(doc, v: O.View, ck: Checker, work: Path, tag: str):
    """Value column as rendered (conditional number format adds "*" when stale)."""
    import subprocess
    if not shutil.which("pdftotext"):
        return
    pdf = doc.export_pdf(work / f"{tag}.pdf")
    txt = subprocess.run(["pdftotext", "-layout", "-f", "1", "-l", "1", str(pdf), "-"],
                         capture_output=True, text=True, check=True).stdout
    for k, o in enumerate(v.sorted[:O.TABLE_ROWS]):
        want = O.rendered_value(o).replace(" ", "")
        m = re.search(rf"(?m)^\s*(?:\S+\s+)?{k + 1}\s+{re.escape(o.name)}\s+(.*?)\s*●", txt)
        ck.eq(f"rendered value row {k + 1} {o.name}", m.group(1).replace(" ", "") if m else None, want)


def shim_vb_object_error(doc):
    """LibreOffice Basic has no VBA constant vbObjectError (Excel: -2147221504), so every
    ``Err.Raise vbObjectError + n`` fails with "Variable not defined" there.  Define it in
    the working copy's modImport so the error texts can be checked."""
    lib = doc.doc.BasicLibraries.getByName(doc.find_library("modImport"))
    code = lib.getByName("modImport")
    if "Const vbObjectError" not in code:
        code = code.replace("Option Explicit\n", "Option Explicit\n"
                            "Private Const vbObjectError As Long = -2147221504\n", 1)
        lib.replaceByName("modImport", code)


def simulate_today(doc, L, day: dt.date | None):
    """Replace c_Today (=TODAY()) by a constant date in the working copy."""
    c = doc.cell("Расчет", L["calc"]["Today"])
    if day is None:
        c.setFormula("=TODAY()")
    else:
        c.setValue(serial(day))
    doc.recalc()


def check_ensure_month(doc, L, db: O.Db, today: dt.date, ck: Checker, start_label: str):
    """modApp.EnsureMonthSelected (Workbook_Open): the current month if it has data,
    otherwise the newest month."""
    set_month(doc, L, start_label)
    doc.run_macro("modApp", "EnsureMonthSelected")
    doc.recalc()
    labels = [lab for _m, lab in O.month_list(db)]
    want = O.month_label(today) if O.month_label(today) in labels else (labels[0] if labels else start_label)
    ck.eq(f"EnsureMonthSelected from {start_label}", doc.get_string("Сводка", L["dash"]["sel"]), want)


# ---------------------------------------------------------------------------
# Synthetic source files
# ---------------------------------------------------------------------------

def _date_cell(ws, r, c, d):
    ws.cell(r, c, EPOCH_DT + (d - EPOCH)).number_format = "dd.mm.yyyy"


def write_source(path: Path, sheets: list):
    """sheets: [(title, header_kind, groups, names, rows)] where header_kind is
    "merged" (Дата in B4:B5, groups row 4, names row 5) or "single" (Дата and
    names in row 3); rows = [(date, [values...])]."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, kind, groups, names, rows in sheets:
        ws = wb.create_sheet(title)
        if kind is None:
            ws["A2"] = "Показатель"
            ws["B3"] = 1.5
            continue
        if kind == "merged":
            ws.merge_cells("B4:B5")
            ws["B4"] = "Дата"
            hdr, grow = 5, 4
        else:
            ws["B3"] = "Дата"
            hdr, grow = 3, None
        for j, nm in enumerate(names):
            ws.cell(hdr, 3 + j, nm)
            if grow and groups:
                ws.cell(grow, 3 + j, groups[j])
        for i, (d, vals) in enumerate(rows):
            _date_cell(ws, hdr + 1 + i, 2, d)
            for j, x in enumerate(vals):
                if x is not None:
                    ws.cell(hdr + 1 + i, 3 + j, x)
    wb.save(path)
    return path


def edge_sources(work: Path) -> dict:
    D = dt.date
    a = write_source(work / "edge_append.xlsx", [
        ("Декабрь 2025", "merged", ["Нефть тип А", "Новые"], ["СИКН №301", "Новый-1"], [
            (D(2025, 12, 29), [0.4, 1.9]),
            (D(2025, 12, 30), [" 1,25 ", "н/д"]),
            (D(2025, 12, 31), [0.9, "-"]),
            (D(2025, 12, 31), [0.45, 2.6]),          # duplicate date row: last wins
        ]),
        ("Свод", None, None, None, None),
        # other column order, other spelling of an existing object, a value that
        # equals the base and one that differs from it
        ("Январь", "single", None, ["Новый-1", "сикн  №304", "СИКН №301"], [
            (D(2026, 1, 1), [0.3, "__SAME304__", "__DIFF301__"]),
            (D(2026, 1, 2), [0.35, None, None]),
        ]),
    ])
    b = write_source(work / "edge_small.xlsx", [
        ("Август", "merged", ["Г1", "Г1", "Г2"], ["Узел-А", "Узел-Б", "Узел-В"], [
            (D(2026, 8, 10), [0.5, 1.6, 2.4]),
            (D(2026, 8, 11), [0.6, 1.7, 3.1]),
            (D(2026, 8, 12), [0.7, None, 0.2]),
        ]),
    ])
    c = write_source(work / "edge_novalues.xlsx", [
        ("Ноябрь", "merged", ["Г"], ["Объект"], [(D(2026, 11, 1), [None]), (D(2026, 11, 2), ["н/д"])]),
    ])
    # 2-decimal rounding at the thresholds, changes between displayed values, near-ties
    rnd = [("СИКН №501", 1.50, 1.504), ("СИКН №502", 1.50, 1.506), ("СИКН №503", 2.30, 2.304),
           ("СИКН №504", 2.30, 2.306), ("СИКН №505", 3.00, 3.004), ("СИКН №506", 3.00, 3.006),
           ("СИКН №507", 1.206, 1.204), ("СИКН №508", 1.20, 1.204), ("СИКН №509", None, 1.8000001),
           ("СИКН №510", None, 1.8000002), ("СИКН №511", 2.66, 2.675), ("СИКН №512", 0.7, None)]
    r = write_source(work / "edge_round.xlsx", [
        ("Июнь", "single", None, [n for n, _a, _b in rnd], [
            (D(2026, 6, 10), [a for _n, a, _b in rnd]),
            (D(2026, 6, 11), [b for _n, _a, b in rnd]),
        ]),
    ])
    # 42 objects (only 40 reach the dashboard, 24 the table), a long file name
    many = [f"Объект {k:02d}" for k in range(1, 43)]

    def mval(k, day):
        if k == 7:
            return None                      # no data at all
        if k == 5:
            return 0.33 if day == 1 else None  # stale
        if k > 40:
            return 9.5 + (k - 41) * 0.4        # beyond the 40 objects of the dashboard
        return round(0.05 * k, 2) if day == 2 else round(0.04 * k, 2)

    mp = write_source(work / "edge_many_objects_with_a_very_long_file_name_to_check_truncation_2026.xlsx", [
        ("Июль", "single", None, many, [
            (D(2026, 7, 1), [mval(k, 1) for k in range(1, 43)]),
            (D(2026, 7, 2), [mval(k, 2) for k in range(1, 43)]),
        ]),
    ])
    # a very long period (more than 60 months in the list) with empty months
    lg = write_source(work / "edge_long.xlsx", [
        ("2017", "single", None, ["А-1", "А-2"], [(D(2017, 1, 15), [0.4, 0.5])]),
        ("2026", "single", None, ["А-1", "А-2"], [(D(2026, 9, 20), [0.6, 2.5])]),
    ])
    too = write_source(work / "edge_toolong.xlsx", [
        ("2010", "single", None, ["А-1"], [(D(2010, 1, 1), [0.4])]),
        ("2026", "single", None, ["А-1"], [(D(2026, 1, 1), [0.6])]),
    ])
    return {"append": a, "small": b, "novalues": c, "round": r, "many": mp, "long": lg,
            "toolong": too, "structure": write_structure_source(work / "edge_structure.xlsx")}


def write_structure_source(path: Path) -> Path:
    """Header variants: a name merged over both header rows, a merged group, service
    columns («Среднее», «ИТОГО …»), an «Итого» row, "<0,05", "н/д", the same dates on a
    copied sheet, a sheet with names and dates but no values, a sheet without «Дата»."""
    D = dt.date
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Май"
    ws.merge_cells("B4:B5")
    ws["B4"] = "Дата"
    ws.merge_cells("C4:D4")
    ws["C4"] = "Нефть"
    ws["C5"], ws["D5"] = "СИКН №901", "СИКН №902"
    ws["E4"], ws["E5"] = "Газ", "СИКН №904"
    ws.merge_cells("F4:F5")
    ws["F4"] = "СИКН №900"                       # no product group
    ws["G5"] = "Среднее"
    ws["H4"], ws["H5"] = "Нефть", "ИТОГО по СИКН"
    ws["I4"], ws["I5"] = "Вода", "СИКН №903"
    rows = [(D(2026, 5, 1), [0.5, 0.6, 1.1, 0.9, 0.7, 3.1, "<0,05"]),
            (D(2026, 5, 2), [0.55, "н/д", 1.2, 1.0, 0.8, 3.6, 0.06]),
            (D(2026, 5, 3), [0.52, 0.61, None, 1.05, 0.75, 3.2, 0.07])]
    for i, (d, vals) in enumerate(rows):
        _date_cell(ws, 6 + i, 2, d)
        for j, x in enumerate(vals):
            if x is not None:
                ws.cell(6 + i, 3 + j, x)
    ws["B9"] = "Итого"
    for j, x in enumerate([1.57, 1.21, 2.3, 2.95, 2.25, 9.9, 0.18]):
        ws.cell(9, 3 + j, x)
    cp = wb.create_sheet("Май (копия)")
    cp["B3"] = "Дата"
    cp["C3"], cp["D3"], cp["E3"] = "СИКН №901", "сикн №900", "Среднее"
    for i, (d, vals) in enumerate([(D(2026, 5, 3), [0.53, 1.07, 0.8]), (D(2026, 5, 4), [0.58, 1.1, 0.84])]):
        _date_cell(cp, 4 + i, 2, d)
        for j, x in enumerate(vals):
            cp.cell(4 + i, 3 + j, x)
    em = wb.create_sheet("Пусто")
    em.merge_cells("B4:B5")
    em["B4"] = "Дата"
    em["C5"] = "СИКН №901"
    _date_cell(em, 6, 2, D(2026, 5, 5))
    _date_cell(em, 7, 2, D(2026, 5, 6))
    sv = wb.create_sheet("Свод")
    sv["A2"] = "Показатель"
    sv["B3"] = 0.9
    wb.save(path)
    return path


def patch_edge_append(path: Path, base: O.Db):
    """Fill the placeholders with the base value (same) and base + 0.5 (conflict)."""
    wb = openpyxl.load_workbook(path)
    ws = wb["Январь"]
    for row in ws.iter_rows():
        for c in row:
            if c.value == "__SAME304__":
                c.value = base.values[(dt.date(2026, 1, 1), "СИКН №304")]
            elif c.value == "__DIFF301__":
                c.value = round(base.values[(dt.date(2026, 1, 1), "СИКН №301")] + 0.5, 2)
    wb.save(path)


# ---------------------------------------------------------------------------
# Scenario
# ---------------------------------------------------------------------------

def default_sel(db: O.Db, today: dt.date) -> dt.date:
    """Month selected after an import: today's month if today is within the data."""
    return O.month_start(today if db.first <= today <= db.last else db.last)


def run(out_dir: Path | None = None) -> Checker:
    from lo_harness import LibreOffice

    work = Path(out_dir) if out_dir else Path(tempfile.mkdtemp(prefix="lhos_e2e_"))
    work.mkdir(parents=True, exist_ok=True)
    book = work / "e2e_dashboard.xlsm"
    shutil.copy(XLSM, book)
    L = discover_layout(book)
    ck = Checker()
    cache = check_static(XLSM, L, ck)
    today = dt.date.today()  # noqa: DTZ011 - must match TODAY() in the workbook
    src_main = O.read_source(MAIN)
    src_upd = O.read_source(UPD)
    default = O.Cfg()
    log_expect = []   # newest first: (mode text, file, total, added, same, conflicts, objects, period)

    def log_add(mode_text, fname, rep):
        period = f"{O.fmt_date(rep.first)} – {O.fmt_date(rep.last)}"
        log_expect.insert(0, (mode_text, fname, rep.total, rep.added, rep.same, rep.conflicts,
                              rep.obj_count, period))

    def keep_autocalc(what):
        """LibreOffice scopes Application.Calculation to the current document: after the
        source workbook is closed, RestoreApp in an error path cannot switch the dashboard
        back to automatic (Excel: application-wide, so it can).  Later imports then keep
        "manual", and LibreOffice stops evaluating conditional formats.  Re-enable it."""
        if not doc.doc.isAutomaticCalculationEnabled():
            ck.notes.append(f"[{ck.ctx}] {what}: LibreOffice left in manual calculation; re-enabled")
            doc.doc.enableAutomaticCalculation(True)

    def run_import(path, mode):
        res = doc.run_macro("modImport", "ImportFileSilent", (str(path), mode))
        keep_autocalc(Path(path).name)
        doc.recalc()
        return res

    with LibreOffice(base_dir=work) as lo:
        lo.set_config("/org.openoffice.Office.Calc/Formula/Load", OOXMLRecalcMode=0)
        lo.set_config("/org.openoffice.Setup/L10N", ooSetupSystemLocale="ru-RU")
        doc = lo.open(book)
        doc.recalc()

        # 0. shipped file vs LibreOffice, VBA sources ----------------------------
        check_cache_matches_lo(doc, L, cache, ck)
        check_vba_sources(doc, ck)

        # 1. empty state ------------------------------------------------------
        ck.ctx = "empty"
        empty = O.Db()
        check_view(doc, L, O.compute(empty, today), ck)
        check_db(doc, empty, ck)
        check_sys(doc, L, ck, file_name=None, mode_text=None, since=None)
        check_freshness(doc, L, ck, None)
        check_info(doc, L, empty, ck)
        ck.eq("log empty", read_log(doc, L), [])
        ck.ctx = "threshold validation rules"
        check_threshold_rules(doc, L, ck)

        # 2. first import, mode 0 --------------------------------------------
        ck.ctx = "import main mode 0"
        t0 = now_local().replace(microsecond=0)
        res = run_import(MAIN, 0)
        db, rep = O.import_source(None, src_main, O.MODE_AUTO)
        check_report(res, rep, src_main, ck)
        check_db(doc, db, ck)
        check_sys(doc, L, ck, file_name=src_main.file_name, mode_text="Полная загрузка", since=t0)
        check_info(doc, L, db, ck)
        log_add("Полная загрузка", src_main.file_name, rep)
        sel = default_sel(db, today)
        ck.eq("sel_Month after import", doc.get_string("Сводка", L["dash"]["sel"]), O.month_label(sel))
        ck.ctx = f"main, today={today}, sel={O.month_label(sel)}"
        v = O.compute(db, today, sel)
        check_view(doc, L, v, ck)
        check_colors(doc, L, v, ck)
        check_pdf_values(doc, v, ck, work, "oct")
        ck.ctx = "elapsed label and freshness dot"
        check_elapsed(doc, L, ck)
        ck.ctx = "EnsureMonthSelected (main)"
        check_ensure_month(doc, L, db, today, ck, "Март 2026")

        # 3. other months via sel_Month --------------------------------------
        for m in (2, 3, 9, 1, 10):
            ms = dt.date(2026, m, 1)
            set_month(doc, L, O.month_label(ms))
            ck.ctx = f"main, sel={O.month_label(ms)}"
            v = O.compute(db, today, ms)
            check_view(doc, L, v, ck)
            if m == 3:
                check_colors(doc, L, v, ck)
        set_month(doc, L, "Сентябрь 2026")
        ck.ctx = "main, sel=Сентябрь 2026 (rendered)"
        check_pdf_values(doc, O.compute(db, today, dt.date(2026, 9, 1)), ck, work, "sep")
        # an invalid selection falls back to the default month
        set_month(doc, L, "Ноябрь 2026")
        ck.ctx = "main, sel=Ноябрь 2026 (not in list)"
        check_view(doc, L, O.compute(db, today, None), ck)
        set_month(doc, L, O.month_label(today))

        # 3b. other "today" dates (c_Today replaced by a constant in the copy) --
        for fake, selm in ((dt.date(2026, 10, 1), 10), (dt.date(2026, 10, 2), 10),
                           (dt.date(2026, 10, 6), 10), (dt.date(2026, 10, 20), 10),
                           (dt.date(2026, 11, 1), 10), (dt.date(2026, 11, 1), None),
                           (dt.date(2026, 3, 17), 3), (dt.date(2026, 3, 18), 3),
                           (dt.date(2026, 3, 1), 3), (dt.date(2026, 2, 10), 2)):
            simulate_today(doc, L, fake)
            set_month(doc, L, O.month_label(dt.date(2026, selm, 1)) if selm else "")
            ck.ctx = f"main, simulated today={fake}, sel={selm}"
            check_view(doc, L, O.compute(db, fake, dt.date(2026, selm, 1) if selm else None), ck,
                       real_today=False)
        simulate_today(doc, L, None)
        set_month(doc, L, O.month_label(today))

        # 4. thresholds -------------------------------------------------------
        for cfg in (O.Cfg(0.5, 1.0, 2.0), O.Cfg(0.7, 0.9, 1.2), O.Cfg(1.8, 2.5, 2.5),
                    O.Cfg(1.25, 2.35, 3.05)):
            set_cfg(doc, L, cfg)
            for m in (10, 3):
                ms = dt.date(2026, m, 1)
                set_month(doc, L, O.month_label(ms))
                ck.ctx = f"main, cfg={cfg}, sel={O.month_label(ms)}"
                check_view(doc, L, O.compute(db, today, ms, cfg), ck)
        set_cfg(doc, L, default)
        set_month(doc, L, O.month_label(today))

        # 5. append import (mode 2) -------------------------------------------
        ck.ctx = "append update mode 2"
        t0 = now_local().replace(microsecond=0)
        res = run_import(UPD, 2)
        db2, rep2 = O.import_source(db, src_upd, O.MODE_APPEND)
        check_report(res, rep2, src_upd, ck)
        ck.eq("oracle: one conflict (СИКН №301 03.10)", [(c[0], c[1]) for c in rep2.conflict_list],
              [("СИКН №301", dt.date(2026, 10, 3))])
        check_db(doc, db2, ck)
        _n, _g, vals, _d, _s = read_db(doc)
        ck.eq("conflict kept base value", vals.get((dt.date(2026, 10, 3), "СИКН №301")),
              src_main.values[(dt.date(2026, 10, 3), "СИКН №301")])
        check_sys(doc, L, ck, file_name=src_upd.file_name, mode_text="Добавление новых", since=t0)
        check_info(doc, L, db2, ck)
        log_add("Добавление новых", src_upd.file_name, rep2)
        sel2 = default_sel(db2, today)
        ck.eq("sel_Month after append", doc.get_string("Сводка", L["dash"]["sel"]), O.month_label(sel2))
        ck.ctx = f"after append, today={today}"
        check_view(doc, L, O.compute(db2, today, sel2), ck)
        set_month(doc, L, "Март 2026")
        ck.ctx = "after append, sel=Март 2026"
        check_view(doc, L, O.compute(db2, today, dt.date(2026, 3, 1)), ck)
        set_month(doc, L, O.month_label(sel2))

        # 6. mode 0 on a non-empty base appends -------------------------------
        ck.ctx = "update again mode 0"
        res = run_import(UPD, 0)
        db3, rep3 = O.import_source(db2, src_upd, O.MODE_AUTO)
        check_report(res, rep3, src_upd, ck)
        ck.eq("mode 0 on non-empty base adds nothing", rep3.added, 0)
        check_db(doc, db3, ck)
        log_add("Добавление новых", src_upd.file_name, rep3)

        # 7. replace import (mode 1) ------------------------------------------
        ck.ctx = "replace main mode 1"
        t0 = now_local().replace(microsecond=0)
        res = run_import(MAIN, 1)
        db4, rep4 = O.import_source(db3, src_main, O.MODE_REPLACE)
        check_report(res, rep4, src_main, ck)
        check_db(doc, db4, ck)
        _n, _g, vals4, _d, _s = read_db(doc)
        ck.eq("replace: base equals main file",
              (_n, len(vals4)), (src_main.objects, len(src_main.values)))
        check_sys(doc, L, ck, file_name=src_main.file_name, mode_text="Полная загрузка", since=t0)
        check_info(doc, L, db4, ck)
        log_add("Полная загрузка", src_main.file_name, rep4)
        ck.ctx = f"after replace, today={today}"
        sel4 = default_sel(db4, today)
        check_view(doc, L, O.compute(db4, today, sel4), ck)

        # 7b. a blank date row inside «БД» (manual edit) -----------------------
        gap = dt.date(2026, 9, 15)
        doc.cell("БД", f"A{3 + (gap - db4.first).days}").clearContents(1 | 2 | 4 | 16)
        doc.recalc()
        db_gap = db4.without_date(gap)
        for selm in (sel4, dt.date(2026, 9, 1)):
            set_month(doc, L, O.month_label(selm))
            ck.ctx = f"blank БД date {gap}, sel={O.month_label(selm)}"
            check_view(doc, L, O.compute(db_gap, today, selm), ck)
        ck.ctx = f"blank БД date {gap}, append update mode 2"
        res = run_import(UPD, 2)
        db_g2, rep_g2 = O.import_source(db_gap, src_upd, O.MODE_APPEND)
        check_report(res, rep_g2, src_upd, ck)
        check_db(doc, db_g2, ck)
        log_add("Добавление новых", src_upd.file_name, rep_g2)
        ck.ctx = "restore main mode 1"
        res = run_import(MAIN, 1)
        db4, rep4 = O.import_source(db_g2, src_main, O.MODE_REPLACE)
        check_report(res, rep4, src_main, ck)
        check_db(doc, db4, ck)
        log_add("Полная загрузка", src_main.file_name, rep4)

        # 8. errors do not touch the base ------------------------------------
        ck.ctx = "error paths"
        shim_vb_object_error(doc)
        edges = edge_sources(work)
        for path, mode, stage, text in (
                (work / "nope.xlsx", 2, "открытие файла", ""),
                (edges["novalues"], 1, "чтение файла", "В файле не найдено ни одного значения."),
                # LibreOffice reports ThisWorkbook.FullName differently, so the check of the
                # dashboard's own path falls through to the "same name is open" check
                (book, 1, "открытие файла", ("Выбран сам файл дашборда.",
                                             f"Уже открыта другая книга с именем «{book.name}».")),
                (edges["toolong"], 1, "запись в базу",
                 (TOO_LONG_TEXT,))):
            res = norm_report(doc.run_macro("modImport", "ImportFileSilent", (str(path), mode)))
            keep_autocalc(Path(path).name)
            head = f"ERR\nОшибка на этапе «{stage}»:\n"
            texts = text if isinstance(text, tuple) else (text,)
            ck.true(f"{Path(path).name} -> ERR at «{stage}»",
                    res.startswith(head) and any(res[len(head):].startswith(t) for t in texts), res[:300])
            check_db(doc, db4, ck)

        # 8b. edge-case source files -------------------------------------------
        ck.ctx = "edge append (Dec 2025, reordered columns, spelling, duplicates)"
        patch_edge_append(edges["append"], db4)
        src_e = O.read_source(edges["append"])
        res = run_import(edges["append"], 2)
        db5, rep5 = O.import_source(db4, src_e, O.MODE_APPEND)
        check_report(res, rep5, src_e, ck)
        check_db(doc, db5, ck)
        log_add("Добавление новых", src_e.file_name, rep5)
        for selm in (dt.date(2025, 12, 1), dt.date(2026, 1, 1)):
            set_month(doc, L, O.month_label(selm))
            ck.ctx = f"edge append, sel={O.month_label(selm)}"
            check_view(doc, L, O.compute(db5, today, selm), ck)

        ck.ctx = "edge replace with a small file"
        src_s = O.read_source(edges["small"])
        res = run_import(edges["small"], 1)
        db6, rep6 = O.import_source(db5, src_s, O.MODE_REPLACE)
        check_report(res, rep6, src_s, ck)
        check_db(doc, db6, ck)
        log_add("Полная загрузка", src_s.file_name, rep6)
        sel6 = default_sel(db6, today)
        ck.eq("sel_Month after small replace", doc.get_string("Сводка", L["dash"]["sel"]),
              O.month_label(sel6))
        ck.ctx = "edge small, default month"
        check_view(doc, L, O.compute(db6, today, sel6), ck)
        for fake in (dt.date(2026, 8, 11), dt.date(2026, 8, 12), dt.date(2026, 8, 13)):
            simulate_today(doc, L, fake)
            ck.ctx = f"edge small, simulated today={fake}"
            check_view(doc, L, O.compute(db6, fake, dt.date(2026, 8, 1)), ck, real_today=False)
        simulate_today(doc, L, None)
        ck.ctx = "EnsureMonthSelected (small)"
        check_ensure_month(doc, L, db6, today, ck, "")

        # 8c. 2-decimal rounding at the zone thresholds -------------------------
        ck.ctx = "edge rounding"
        src_r = O.read_source(edges["round"])
        res = run_import(edges["round"], 1)
        db7, rep7 = O.import_source(db6, src_r, O.MODE_REPLACE)
        check_report(res, rep7, src_r, ck)
        check_db(doc, db7, ck)
        log_add("Полная загрузка", src_r.file_name, rep7)
        v7 = O.compute(db7, today, default_sel(db7, today))
        ck.eq("oracle: zones 1.504/1.506/2.304/2.306/3.004/3.006",
              [o.zone for o in v7.objects[:6]], [1, 2, 2, 3, 3, 4])
        ck.eq("oracle: arrows (0 change neutral, -0.002 raw -> ↓)",
              [O.arrow_of(o.change) for o in v7.objects[:8]], ["—", "↑", "—", "↑", "—", "↑", "↓", "—"])
        check_view(doc, L, v7, ck)
        check_colors(doc, L, v7, ck)
        check_pdf_values(doc, v7, ck, work, "round")
        for cfg in (O.Cfg(1.25, 2.35, 3.05), O.Cfg(1.51, 2.3, 3.0)):
            set_cfg(doc, L, cfg)
            ck.ctx = f"edge rounding, cfg={cfg}"
            check_view(doc, L, O.compute(db7, today, default_sel(db7, today), cfg), ck)
        set_cfg(doc, L, default)

        # 8d. more objects than the table / the dashboard shows -----------------
        ck.ctx = "edge 42 objects, long file name"
        src_m = O.read_source(edges["many"])
        res = run_import(edges["many"], 1)
        db8, rep8 = O.import_source(db7, src_m, O.MODE_REPLACE)
        check_report(res, rep8, src_m, ck)
        check_db(doc, db8, ck)
        check_info(doc, L, db8, ck)
        check_sys(doc, L, ck, file_name=src_m.file_name, mode_text="Полная загрузка",
                  since=now_local() - dt.timedelta(minutes=1))
        log_add("Полная загрузка", src_m.file_name, rep8)
        v8 = O.compute(db8, today, default_sel(db8, today))
        ck.eq("oracle: note for 42 objects", O.table_note(v8),
              "Показаны 24 из 42 СИКН; * — значение за предыдущую дату")
        check_view(doc, L, v8, ck)

        # 8e. header variants, service columns, dates repeated on two sheets ----
        ck.ctx = "edge structure"
        src_t = O.read_source(edges["structure"])
        ck.eq("oracle: structure objects", [(n, src_t.groups[n]) for n in src_t.objects],
              [("СИКН №901", "Нефть"), ("СИКН №902", "Нефть"), ("СИКН №904", "Газ"),
               ("СИКН №900", ""), ("СИКН №903", "Вода")])
        ck.eq("oracle: duplicates", (src_t.dup_count, src_t.dup_example),
              (2, ("Май", "Май (копия)", dt.date(2026, 5, 3))))
        res = run_import(edges["structure"], 1)
        db9, rep9 = O.import_source(db8, src_t, O.MODE_REPLACE)
        check_report(res, rep9, src_t, ck)
        check_db(doc, db9, ck)
        log_add("Полная загрузка", src_t.file_name, rep9)
        check_view(doc, L, O.compute(db9, today, default_sel(db9, today)), ck)

        # 8f. 9 years of data: month list, months without values ---------------
        ck.ctx = "edge long period"
        src_l = O.read_source(edges["long"])
        res = run_import(edges["long"], 1)
        db10, rep10 = O.import_source(db9, src_l, O.MODE_REPLACE)
        check_report(res, rep10, src_l, ck)
        check_db(doc, db10, ck)
        log_add("Полная загрузка", src_l.file_name, rep10)
        ck.eq("oracle: months in the list", len(O.month_list(db10)), 117)
        sel10 = default_sel(db10, today)
        for selm in (sel10, dt.date(2020, 6, 1), dt.date(2017, 1, 1)):
            set_month(doc, L, O.month_label(selm))
            ck.ctx = f"edge long period, sel={O.month_label(selm)}"
            check_view(doc, L, O.compute(db10, today, selm), ck)
        simulate_today(doc, L, dt.date(2026, 9, 10))
        set_month(doc, L, "Сентябрь 2026")
        ck.ctx = "edge long period, simulated today=2026-09-10 (no values yet)"
        check_view(doc, L, O.compute(db10, dt.date(2026, 9, 10), dt.date(2026, 9, 1)), ck,
                   real_today=False)
        simulate_today(doc, L, None)

        # 9. clear -------------------------------------------------------------
        ck.ctx = "clear"
        res = doc.run_macro("modImport", "ClearDatabaseSilent")
        keep_autocalc("ClearDatabaseSilent")
        doc.recalc()
        ck.eq("ClearDatabaseSilent returns True", res, True)
        check_db(doc, O.Db(), ck)
        check_view(doc, L, O.compute(O.Db(), today), ck)
        check_sys(doc, L, ck, file_name=None, mode_text=None, since=None)
        check_freshness(doc, L, ck, None)
        check_info(doc, L, O.Db(), ck)
        ck.eq("sel_Month cleared", doc.get_string("Сводка", L["dash"]["sel"]), "")
        log_expect.insert(0, ("Очистка базы", None, 0, 0, 0, 0, 0, None))

        # 10. import log -------------------------------------------------------
        ck.ctx = "import log"
        log = read_log(doc, L)
        ck.eq("log row count", len(log), len(log_expect))
        for i, (want, got) in enumerate(zip(log_expect, log)):
            mode_text, fname, total, added, same, conf, objs, period = want
            ck.eq(f"log[{i}] mode", got[2], mode_text)
            ck.eq(f"log[{i}] file", got[1], fname)
            ck.eq(f"log[{i}] numbers", got[3:8], [float(x) for x in (total, added, same, conf, objs)])
            ck.eq(f"log[{i}] period", got[8], period)
            ck.true(f"log[{i}] timestamp", isinstance(got[0], float), str(got[0]))
        doc.close()

    if out_dir:
        (work / "mismatches.json").write_text(
            json.dumps(ck.fails, ensure_ascii=False, indent=1, default=str), encoding="utf-8")
        (work / "notes.json").write_text(json.dumps(ck.notes, ensure_ascii=False, indent=1),
                                         encoding="utf-8")
    else:
        shutil.rmtree(work, ignore_errors=True)
    return ck


def summarize(ck: Checker, limit: int = 400) -> str:
    lines = [f"{ck.count} checks, {len(ck.fails)} mismatches"]
    lines += [f"note: {n}" for n in ck.notes]
    for f in ck.fails[:limit]:
        lines.append(f"[{f['ctx']}] {f['what']}: actual={f['actual']!r} expected={f['expected']!r}")
    return "\n".join(lines)


def test_integration():
    out = os.environ.get("LHOS_E2E_OUT")
    ck = run(Path(out) if out else None)
    assert not ck.fails, summarize(ck)


if __name__ == "__main__":
    if not HAVE_LO:
        print("SKIP: LibreOffice/uno not available")
        sys.exit(0)
    t = time.time()
    out = os.environ.get("LHOS_E2E_OUT")
    ck = run(Path(out) if out else None)
    print(summarize(ck))
    print(f"{time.time() - t:.1f} s")
    sys.exit(1 if ck.fails else 0)
