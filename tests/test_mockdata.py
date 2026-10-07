"""Tests for src/lhos/mockdata.py (mock source workbooks).

Runs under pytest or standalone: ``python3 -I tests/test_mockdata.py``
(exit code 1 on failure). Set ``LHOS_SOFFICE_TESTS=1`` to also recalculate the
"Свод" formulas in headless LibreOffice and compare them with Python averages.
"""

from __future__ import annotations

import atexit
import datetime as dt
import functools
import os
import shutil
import statistics
import subprocess
import sys
import tempfile
import traceback
import zipfile
from pathlib import Path
from xml.etree import ElementTree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

import openpyxl  # noqa: E402

from lhos import mockdata as md  # noqa: E402

MAIN_NAME = "Мокап_ЛХОС_2026.xlsx"
UPDATE_NAME = "Мокап_ЛХОС_2026_дополнение.xlsx"
EXPECTED_MERGES = {"B4:B5", "C4:E4", "F4:H4", "K4:R4", "S4:T4", "U4:V4"}
EXPECTED_GROUP_CELLS = {
    "C4": "Нефть тип А", "F4": "Нефть тип Б", "I4": "Нефть тип Б", "J4": "Нефть тип Б",
    "K4": "Нефть тип А", "S4": "Нефть тип Б", "U4": "Узлы подготовки",
}
EXPECTED_NAMES = [
    "СИКН №301", "СИКН №304", "СИКН №307", "СИКН №310", "СИКН №312", "СИКН №315",
    "СИКН №318", "ПСП Восточный", "СИКН №320", "СИКН №322", "СИКН №325", "СИКН №327",
    "СИКН №330", "СИКН №333", "СИКН №336", "СИКН №339", "СИКН №341", "СИКН №344",
    "УПСВ-1 Лесная", "УПСВ-2 Озерная",
]
MONTH_SHEETS = list(md.MONTHS_RU[:11])


# ---------------------------------------------------------------------------
# Fixtures / helpers
# ---------------------------------------------------------------------------

@functools.lru_cache(maxsize=None)
def generated() -> tuple[Path, dict]:
    """Generate both files once into a temporary directory (removed at exit)."""
    out = Path(tempfile.mkdtemp(prefix="lhos_mock_"))
    atexit.register(shutil.rmtree, out, ignore_errors=True)
    summary = md.make_mock(str(out / MAIN_NAME), str(out / UPDATE_NAME), seed=42)
    return out, summary


@functools.lru_cache(maxsize=None)
def workbook(kind: str):
    out, _ = generated()
    return openpyxl.load_workbook(out / (MAIN_NAME if kind == "main" else UPDATE_NAME))


def iso(value) -> str:
    return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()


@functools.lru_cache(maxsize=None)
def read_values(kind: str) -> dict[tuple[str, str], object]:
    """{(ISO date, object name): value} for every non-empty value cell, read like an importer."""
    wb = workbook(kind)
    values = {}
    for sheet in MONTH_SHEETS:
        ws = wb[sheet]
        names = {}
        col = 3
        while ws.cell(5, col).value:
            names[col] = ws.cell(5, col).value
            col += 1
        row = 6
        while ws.cell(row, 2).value is not None:
            day = iso(ws.cell(row, 2).value)
            for col, name in names.items():
                value = ws.cell(row, col).value
                if value is not None:
                    values[(day, name)] = value
            row += 1
    return values


def is_number(value) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def month_days(month: int) -> int:
    return (dt.date(2026 + month // 12, month % 12 + 1, 1) - dt.date(2026, month, 1)).days


# ---------------------------------------------------------------------------
# Layout
# ---------------------------------------------------------------------------

def test_sheet_order():
    for kind in ("main", "update"):
        assert workbook(kind).sheetnames == MONTH_SHEETS + ["Свод"], kind
        assert workbook(kind).active.title == "Октябрь"


def test_month_sheet_layout():
    for kind in ("main", "update"):
        for month, sheet in enumerate(MONTH_SHEETS, start=1):
            ws = workbook(kind)[sheet]
            where = f"{kind}/{sheet}"
            assert {str(r) for r in ws.merged_cells.ranges} == EXPECTED_MERGES, where
            assert ws["B4"].value == "Дата", where
            for coord, text in EXPECTED_GROUP_CELLS.items():
                assert ws[coord].value == text, (where, coord)
            extra = kind == "update" and sheet == "Октябрь"
            names = [ws.cell(5, c).value for c in range(3, 3 + 20 + (1 if extra else 0))]
            assert names == EXPECTED_NAMES + (["СИКН №350"] if extra else []), where
            assert ws.cell(5, 3 + len(names)).value is None, where
            for c in range(3, 3 + len(names)):
                cell = ws.cell(5, c)
                assert cell.font.name == "Arial Cyr" and cell.font.sz == 8, (where, cell.coordinate)
                assert cell.alignment.horizontal == "center" and cell.alignment.wrap_text
                assert all(getattr(cell.border, s).style == "thin"
                           for s in ("left", "right", "top", "bottom"))
            assert abs(ws.column_dimensions["B"].width - 11.71) < 0.01, where
            assert ws.sheet_view.zoomScale == 85, where
            assert len(ws._charts) == 1, where

            # Dates: every day of the month from B6, builtin format 14, nothing after.
            days = month_days(month)
            for day in range(1, days + 1):
                cell = ws.cell(5 + day, 2)
                assert iso(cell.value) == dt.date(2026, month, day).isoformat(), (where, day)
                assert cell.number_format == "mm-dd-yy", (where, cell.number_format)
                assert cell.alignment.horizontal == "center"
                assert cell.border.left.style == "thin" and cell.border.bottom.style == "thin"
                assert ws.cell(5 + day, 3).border.top.style == "thin"   # empty cells framed too
            assert ws.cell(6 + days, 2).value is None, where


def raw_cell_format(path: Path, sheet_xml: str, coord: str) -> dict:
    """Font and border of a cell straight from the package XML.

    openpyxl drops the formats of cells covered by a merge when it loads a file,
    so covered header cells are checked at the XML level.
    """
    ns = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    with zipfile.ZipFile(path) as z:
        styles = ElementTree.fromstring(z.read("xl/styles.xml"))
        sheet = ElementTree.fromstring(z.read(f"xl/worksheets/{sheet_xml}"))
    cell = sheet.find(f".//m:c[@r='{coord}']", ns)
    xf = styles.findall("m:cellXfs/m:xf", ns)[int(cell.get("s", 0))]
    font = styles.findall("m:fonts/m:font", ns)[int(xf.get("fontId"))]
    border = styles.findall("m:borders/m:border", ns)[int(xf.get("borderId"))]
    return {
        "font": (font.find("m:name", ns).get("val"), float(font.find("m:sz", ns).get("val"))),
        "border": {side: border.find(f"m:{side}", ns).get("style")
                   for side in ("left", "right", "top", "bottom")},
    }


def test_group_header_styles_mirror_template():
    out, _ = generated()
    path = out / MAIN_NAME          # sheet1.xml = "Январь"
    assert raw_cell_format(path, "sheet1.xml", "L4")["font"] == ("Arial Cyr", 10.0)
    assert raw_cell_format(path, "sheet1.xml", "D4")["font"] == ("Calibri", 11.0)
    assert raw_cell_format(path, "sheet1.xml", "C5")["font"] == ("Arial Cyr", 8.0)
    assert raw_cell_format(path, "sheet1.xml", "U4")["border"] == {
        "left": "thin", "right": None, "top": "thin", "bottom": "thin"}
    assert raw_cell_format(path, "sheet1.xml", "V4")["border"] == {
        "left": None, "right": "thin", "top": "thin", "bottom": "thin"}


def test_every_header_cell_matches_template():
    """Row 4/5 of a month sheet and row 2/3 of "Свод", cell by cell (raw XML).

    Expected formats were read from the user's template: covered cells of the
    K:R and S:T merges use Arial Cyr 10, U:V is drawn as two half-open boxes,
    everything else in the group row is a Calibri 11 thin box.
    """
    out, _ = generated()
    path = out / MAIN_NAME          # sheet1.xml = "Январь", sheet12.xml = "Свод"
    thin = {"left": "thin", "right": "thin", "top": "thin", "bottom": "thin"}
    for c in range(3, 23):
        month_col = openpyxl.utils.get_column_letter(c)
        svod_col = openpyxl.utils.get_column_letter(c - 1)
        font = ("Arial Cyr", 10.0) if month_col in "LMNOPQRT" else ("Calibri", 11.0)
        border = dict(thin)
        if month_col == "U":
            border["right"] = None
        elif month_col == "V":
            border["left"] = None
        for sheet_xml, group_cell, name_cell in (("sheet1.xml", f"{month_col}4", f"{month_col}5"),
                                                 ("sheet12.xml", f"{svod_col}2", f"{svod_col}3")):
            got = raw_cell_format(path, sheet_xml, group_cell)
            assert got == {"font": font, "border": border}, (sheet_xml, group_cell, got)
            got = raw_cell_format(path, sheet_xml, name_cell)
            assert got == {"font": ("Arial Cyr", 8.0), "border": thin}, (sheet_xml, name_cell, got)


def test_defined_names():
    wb = workbook("main")
    names = {n: wb.defined_names[n].attr_text for n in wb.defined_names}
    assert len(names) == 20
    assert names["СИКН_№301"] == "Январь!$C$6:$C$36"
    assert names["УПСВ_2_Озерная"] == "Январь!$V$6:$V$36"


def test_svod_sheet():
    for kind in ("main", "update"):
        ws = workbook(kind)["Свод"]
        merges = {str(r) for r in ws.merged_cells.ranges}
        assert merges == {"A2:A3", "B2:D2", "E2:G2", "J2:Q2", "R2:S2", "T2:U2"}, merges
        assert ws["A2"].value == "Показатель"
        assert ws["A4"].value == "Среднее значение за месяц:"
        assert ws["A17"].value == "Среднее значение за год"
        assert [ws.cell(3, c).value for c in range(2, 22)] == EXPECTED_NAMES
        assert [ws.cell(r, 1).value for r in range(5, 17)] == list(md.MONTHS_RU)
        for c in range(2, 22):
            month_col = openpyxl.utils.get_column_letter(c + 1)
            for month in range(1, 12):
                last = 5 + month_days(month)
                assert ws.cell(4 + month, c).value == (
                    f"=AVERAGE({md.MONTHS_RU[month - 1]}!{month_col}$6:{month_col}${last})")
                assert ws.cell(4 + month, c).number_format == "0.00"
            assert ws.cell(16, c).value is None          # no "Декабрь" sheet yet
            year = ws.cell(17, c).value
            assert year.startswith(f"=AVERAGE(Январь!{month_col}6:{month_col}36,")
            assert year.count("!") == 11 and "Ноябрь" in year
        assert ws["T5"].border.right.style == "medium" and ws["A5"].border.left.style == "medium"
        rules = [(str(cf.sqref), rule) for cf in ws.conditional_formatting for rule in cf.rules]
        assert len(rules) == 1
        sqref, rule = rules[0]
        assert sqref == "B5:U17" and rule.type == "cellIs" and rule.operator == "greaterThan"
        assert rule.formula == ["1"]
        assert rule.dxf.fill.bgColor.rgb == "FFFFFF00"


def test_package_like_excel_and_deterministic():
    out, _ = generated()
    for name in (MAIN_NAME, UPDATE_NAME):
        with zipfile.ZipFile(out / name) as z:
            assert "xl/sharedStrings.xml" in z.namelist()
            for part in z.namelist():
                if part.startswith("xl/worksheets/sheet"):
                    assert b"inlineStr" not in z.read(part), part
    with tempfile.TemporaryDirectory() as tmp:
        again = md.make_mock(str(Path(tmp) / MAIN_NAME), str(Path(tmp) / UPDATE_NAME), seed=42)
        for name in (MAIN_NAME, UPDATE_NAME):
            assert (Path(tmp) / name).read_bytes() == (out / name).read_bytes(), name
        assert again["special"] == generated()[1]["special"]
        other = md.make_mock(str(Path(tmp) / MAIN_NAME), str(Path(tmp) / UPDATE_NAME), seed=7)
        assert other["main"]["latest"] != generated()[1]["main"]["latest"]


# ---------------------------------------------------------------------------
# Values
# ---------------------------------------------------------------------------

def test_values_match_summary():
    _, summary = generated()
    for kind in ("main", "update"):
        values = read_values(kind)
        info = summary[kind]
        assert sum(map(is_number, values.values())) == info["numeric_values"], kind
        assert sum(isinstance(v, str) for v in values.values()) == info["text_values"] == 2
        assert min(d for d, _ in values) == info["date_first"] == "2026-01-01"
        assert max(d for d, _ in values) == info["date_last"]
        assert all(v >= 0 for v in values.values() if is_number(v))
        assert not any(d.startswith("2026-11") for d, _ in values), "Ноябрь must stay empty"
    assert summary["main"]["date_last"] == "2026-10-06"
    assert summary["update"]["date_last"] == "2026-10-07"
    assert [o["name"] for o in summary["objects"]] == EXPECTED_NAMES
    assert [o["col"] for o in summary["objects"]][0] == "C"


def test_realism_ratios():
    values = read_values("main")
    _, summary = generated()
    cells = 279 * 20                       # 2026-01-01 .. 2026-10-06, 20 objects
    blanks = cells - len(values)
    assert blanks == len(summary["main"]["blank_cells"])
    assert 0.015 <= blanks / cells <= 0.03, blanks
    numbers = [v for v in values.values() if is_number(v)]
    two_dec = sum(round(v, 1) != v for v in numbers) / len(numbers)
    assert 0.1 <= two_dec <= 0.3, two_dec


def test_special_cells():
    _, summary = generated()
    special = summary["special"]
    values = read_values("main")
    wb = workbook("main")

    def at(ref):
        return wb[ref["sheet"]][ref["cell"]].value

    # Exactly one exceedance of the 3.0 limit.
    over = [(k, v) for k, v in values.items() if is_number(v) and v > md.LIMIT]
    assert over == [(("2026-03-17", "СИКН №327"), 3.2)], over
    assert at(special["exceedance"]) == 3.2
    assert special["exceedance"]["sheet"] == "Март" and special["exceedance"]["cell"] == "N22"

    # Operator-typed text.
    assert values[("2026-02-10", "СИКН №312")] == "н/д" == at(special["text_value"])
    assert values[("2026-03-03", "СИКН №330")] == "0,45" == at(special["numeric_text"])
    assert {v for v in values.values() if isinstance(v, str)} == {"н/д", "0,45"}

    # Yellow and red days.
    assert 3 <= len(special["yellow_days"]) and 3 <= len(special["red_days"])
    for ref in special["yellow_days"]:
        assert 1.6 <= at(ref) <= 2.2 and ref["date"] < "2026-10-01"
    for ref in special["red_days"]:
        assert 2.4 <= at(ref) <= 2.9 and ref["date"] < "2026-10-01"

    # Latest day of the main file.
    day = "2026-10-06"
    assert values[(day, "СИКН №322")] == 1.8
    assert values[(day, "СИКН №307")] == 2.5
    assert (day, "СИКН №341") not in values
    assert values[("2026-10-05", "СИКН №341")] == 0.6
    for name in EXPECTED_NAMES:
        if name not in ("СИКН №322", "СИКН №307", "СИКН №341"):
            assert values[(day, name)] <= md.GREEN_MAX, name
    october = [(k, v) for k, v in values.items() if k[0].startswith("2026-10")]
    assert {k for k, v in october if v > md.GREEN_MAX} == {(day, "СИКН №322"), (day, "СИКН №307")}

    # Non-green cells carry the operator's yellow highlight, green ones do not.
    ws = wb["Октябрь"]
    assert ws["L11"].fill.fgColor.rgb == "FFFFFF00" and ws["E11"].fill.fill_type == "solid"
    assert ws["C11"].fill.fill_type is None

    # Month-end blank -> previous day.
    blank = special["month_end_blank"]
    assert at(blank["blank"]) is None and at(blank["previous"]) == blank["previous"]["value"]


def test_latest_and_month_last():
    _, summary = generated()
    for kind in ("main", "update"):
        values = read_values(kind)
        names = EXPECTED_NAMES + (["СИКН №350"] if kind == "update" else [])
        latest = {}
        month_last = {}
        for (day, name), value in sorted(values.items()):
            if not is_number(value):
                continue
            latest[name] = {"date": day, "value": value}
            month = md.MONTHS_RU[int(day[5:7]) - 1]
            month_last.setdefault(month, {})[name] = {"date": day, "value": value}
        assert latest == summary[kind]["latest"], kind
        assert set(latest) == set(names)
        assert month_last == summary[kind]["month_last"], kind
    assert summary["main"]["latest"]["СИКН №341"] == {"date": "2026-10-05", "value": 0.6}
    assert summary["main"]["month_last"]["Сентябрь"]["СИКН №315"]["date"] == "2026-09-29"


def test_update_differs_only_as_specified():
    main, update = read_values("main"), read_values("update")
    _, summary = generated()
    changed = summary["special"]["update_changed"]
    changed_key = (changed["date"], changed["object"])
    assert abs(update[changed_key] - main[changed_key] - 0.3) < 1e-9
    assert changed["original"] == main[changed_key] and changed["value"] == update[changed_key]

    added = {k: v for k, v in update.items() if k not in main}
    new_day = {k for k in added if k[0] == "2026-10-07"}
    new_obj = {k for k in added if k[1] == "СИКН №350"}
    assert added.keys() == new_day | new_obj
    assert {n for _, n in new_day} == set(EXPECTED_NAMES) | {"СИКН №350"}
    assert sorted(d for d, _ in new_obj) == [f"2026-10-0{i}" for i in range(1, 8)]
    assert all(v <= md.GREEN_MAX for v in added.values())
    assert not set(main) - set(update), "update must not drop values"
    diffs = [k for k in main if main[k] != update[k]]
    assert diffs == [changed_key], diffs

    ws = workbook("update")["Октябрь"]
    assert ws["W4"].value == "Нефть тип А" and ws["W5"].value == "СИКН №350"
    assert ws["W4"].border.left.style == "thin"
    for sheet in MONTH_SHEETS:
        if sheet != "Октябрь":
            assert workbook("update")[sheet]["W5"].value is None, sheet
    assert workbook("main")["Октябрь"]["W5"].value is None


def test_libreoffice_recalculation():
    """Opt-in: LibreOffice evaluates the "Свод" formulas to the Python averages."""
    soffice = shutil.which("soffice")
    if os.environ.get("LHOS_SOFFICE_TESTS") != "1" or not soffice:
        print("  (skipped: set LHOS_SOFFICE_TESTS=1 and install soffice)")
        return
    out, _ = generated()
    with tempfile.TemporaryDirectory() as tmp:
        src = Path(tmp) / "main.xlsx"
        shutil.copy(out / MAIN_NAME, src)
        subprocess.run([soffice, f"-env:UserInstallation=file://{tmp}/profile", "--headless",
                        "--convert-to", "xlsx", "--outdir", f"{tmp}/out", str(src)],
                       check=True, timeout=180, capture_output=True)
        ws = openpyxl.load_workbook(Path(tmp) / "out" / "main.xlsx", data_only=True)["Свод"]
    values = read_values("main")
    for c, name in enumerate(EXPECTED_NAMES, start=2):
        for month in range(1, 11):
            nums = [v for (d, n), v in values.items()
                    if n == name and int(d[5:7]) == month and is_number(v)]
            assert abs(ws.cell(4 + month, c).value - statistics.fmean(nums)) < 1e-9, (name, month)
        assert ws.cell(15, c).value == "#DIV/0!"        # empty "Ноябрь", as in the template
        year = [v for (d, n), v in values.items() if n == name and is_number(v)]
        assert abs(ws.cell(17, c).value - statistics.fmean(year)) < 1e-9, name


# ---------------------------------------------------------------------------

def main() -> int:
    tests = [(n, f) for n, f in globals().items() if n.startswith("test_") and callable(f)]
    failed = 0
    for name, func in tests:
        try:
            func()
            print(f"PASS {name}")
        except Exception:
            failed += 1
            print(f"FAIL {name}")
            traceback.print_exc()
    print(f"{len(tests) - failed}/{len(tests)} passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
