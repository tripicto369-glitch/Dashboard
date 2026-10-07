"""Independent reference model ("oracle") of the LHOS dashboard.

Re-implements, in plain Python and *without* looking at the workbook's
formulas, what the dashboard must show:

* reading a source workbook in the user's format (one sheet per month, a
  "Дата" header cell, object names in the bottom header row, product groups in
  the row above, dates below, values in ppm; "0,45" typed as text is a number,
  "н/д" is skipped; sheets without "Дата" are ignored);
* the database after an import: first import fills it, later imports either
  replace everything (mode 1) or only add values that are not in the base yet
  (mode 2; differing values are conflicts and the base keeps its value);
* the dashboard for a given TODAY and selected month:
    - current month: latest value on/before TODAY (if TODAY has none, the
      previous day(s) of the month; on the 1st the previous day may be in the
      previous month);
    - past month: last value of that month;
    - change (24h) = value minus the previous measured value of the object;
    - zones: v <= green -> 1, <= yellow -> 2, <= limit -> 3, > limit -> 4;
    - table sorted by value (desc), ties in base order, objects without value
      last; "*" when an object's value is older than the newest value date;
    - KPIs (max + object, average, zone counts, total, overall state);
    - trend: daily values of the selected month up to the reference date;
    - max-per-day series for the same days;
    - month list: months from the first to the last month with data, newest
      first.

Only the standard library and openpyxl are used.
"""

from __future__ import annotations

import calendar
import datetime as dt
import math
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from typing import Optional

import openpyxl

MONTHS_RU = ("Январь", "Февраль", "Март", "Апрель", "Май", "Июнь", "Июль", "Август",
             "Сентябрь", "Октябрь", "Ноябрь", "Декабрь")
MAX_OBJ = 40
TABLE_ROWS = 24

MODE_AUTO, MODE_REPLACE, MODE_APPEND = 0, 1, 2

STATE_TEXT = ("НЕТ ДАННЫХ", "НОРМА", "ВНИМАНИЕ", "РИСК", "ПРЕВЫШЕНИЕ")
ROW_STATUS = ("●  НЕТ ДАННЫХ", "●  НОРМА", "●  ВНИМАНИЕ", "●  РИСК", "●  ПРЕВЫШЕНИЕ")


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def clean(v) -> str:
    if v is None:
        return ""
    s = str(v).replace("\r", " ").replace("\n", " ").replace("\t", " ").replace("\xa0", " ")
    return " ".join(s.split())


def as_date(v) -> Optional[dt.date]:
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    return None


def parse_number(v):
    """(kind, value): kind 1 numeric, 0 empty, -1 non-numeric text."""
    if v is None:
        return 0, None
    if isinstance(v, bool):
        return -1, None
    if isinstance(v, (int, float)):
        return 1, float(v)
    s = clean(v).replace(" ", "")
    if not s:
        return 0, None
    s = s.replace(",", ".")
    try:
        x = float(s)
    except ValueError:
        return -1, None
    if math.isnan(x) or math.isinf(x):
        return -1, None
    return 1, x


def month_start(d: dt.date) -> dt.date:
    return d.replace(day=1)


def month_end(d: dt.date) -> dt.date:
    return d.replace(day=calendar.monthrange(d.year, d.month)[1])


def month_label(d: dt.date) -> str:
    return f"{MONTHS_RU[d.month - 1]} {d.year}"


def add_months(d: dt.date, n: int) -> dt.date:
    y, m = divmod(d.year * 12 + d.month - 1 + n, 12)
    return dt.date(y, m + 1, 1)


def fmt_num(x: float, min_dec: int = 1, max_dec: int = 2, comma: str = ",") -> str:
    """Excel number format like ``0.0#`` (round half away from zero)."""
    q = Decimal(repr(abs(x))).quantize(Decimal(1).scaleb(-max_dec), rounding=ROUND_HALF_UP)
    s = f"{q:.{max_dec}f}"
    ip, fp = s.split(".")
    while len(fp) > min_dec and fp.endswith("0"):
        fp = fp[:-1]
    out = ip + ((comma + fp) if fp else "")
    neg = x < 0 and q != 0
    return ("-" if neg else "") + out


def fmt_change(x: float, comma: str = ",") -> str:
    """Excel number format ``+0.0#;-0.0#;0.0``."""
    q = Decimal(repr(abs(x))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    if x > 0 and q != 0:
        return "+" + fmt_num(x, comma=comma)
    if x < 0 and q != 0:
        return "-" + fmt_num(-x, comma=comma)
    # Excel picks the section by the *unrounded* sign
    if x > 0:
        return "+" + fmt_num(x, comma=comma)
    if x < 0:
        return "-" + fmt_num(-x, comma=comma)
    return "0" + comma + "0"


def fmt_date(d: Optional[dt.date]) -> str:
    return d.strftime("%d.%m.%Y") if d else "—"


# ---------------------------------------------------------------------------
# Source file
# ---------------------------------------------------------------------------

@dataclass
class Source:
    file_name: str
    objects: list = field(default_factory=list)          # first-seen order
    groups: dict = field(default_factory=dict)
    values: dict = field(default_factory=dict)           # (date, name) -> float; last wins
    sheets_with_values: list = field(default_factory=list)
    empty_sheets: int = 0
    text_skipped: int = 0

    @property
    def total(self) -> int:
        return len(self.values)

    @property
    def first(self) -> dt.date:
        return min(d for d, _ in self.values)

    @property
    def last(self) -> dt.date:
        return max(d for d, _ in self.values)


def _merged_top_left(ws, row: int, col: int):
    for m in ws.merged_cells.ranges:
        if m.min_row <= row <= m.max_row and m.min_col <= col <= m.max_col:
            return ws.cell(m.min_row, m.min_col).value, m
    return ws.cell(row, col).value, None


def read_source(path) -> Source:
    import os
    wb = openpyxl.load_workbook(path, data_only=True)
    src = Source(os.path.basename(str(path)))
    keys: dict[str, str] = {}   # upper-case key -> canonical name
    for ws in wb.worksheets:
        hdr = None
        for row in ws.iter_rows(min_row=1, max_row=40, max_col=26):
            for c in row:
                if isinstance(c.value, str) and clean(c.value).lower() in ("дата", "дата:"):
                    hdr = c
                    break
            if hdr:
                break
        if hdr is None:
            continue
        _v, merge = _merged_top_left(ws, hdr.row, hdr.column)
        if merge is not None:
            name_row, group_row = merge.max_row, (merge.min_row if merge.max_row > merge.min_row else None)
        else:
            name_row, group_row = hdr.row, None
        date_col = hdr.column
        cols = {}
        for col in range(date_col + 1, ws.max_column + 1):
            nm = clean(ws.cell(name_row, col).value)
            if not nm:
                continue
            grp = clean(_merged_top_left(ws, group_row, col)[0]) if group_row else ""
            key = nm.upper()
            if key not in keys:
                keys[key] = nm
                src.objects.append(nm)
                src.groups[nm] = grp
            elif not src.groups.get(keys[key]):
                src.groups[keys[key]] = grp
            cols[col] = keys[key]
        n = 0
        for r in range(name_row + 1, ws.max_row + 1):
            d = as_date(ws.cell(r, date_col).value)
            if d is None:
                continue
            for col, nm in cols.items():
                kind, x = parse_number(ws.cell(r, col).value)
                if kind == 1:
                    src.values[(d, nm)] = x
                    n += 1
                elif kind == -1:
                    src.text_skipped += 1
        if n:
            src.sheets_with_values.append(ws.title)
        else:
            src.empty_sheets += 1
    return src


# ---------------------------------------------------------------------------
# Database and import
# ---------------------------------------------------------------------------

@dataclass
class Db:
    objects: list = field(default_factory=list)
    groups: dict = field(default_factory=dict)
    values: dict = field(default_factory=dict)

    def copy(self) -> "Db":
        return Db(list(self.objects), dict(self.groups), dict(self.values))

    @property
    def empty(self) -> bool:
        return not self.values

    @property
    def first(self) -> Optional[dt.date]:
        return min((d for d, _ in self.values), default=None)

    @property
    def last(self) -> Optional[dt.date]:
        return max((d for d, _ in self.values), default=None)

    def dates(self) -> list:
        if self.empty:
            return []
        n = (self.last - self.first).days + 1
        return [self.first + dt.timedelta(days=i) for i in range(n)]

    def get(self, d: dt.date, name: str):
        return self.values.get((d, name))


@dataclass
class ImportReport:
    mode: int
    total: int
    added: int
    same: int
    conflicts: int
    conflict_list: list
    new_objects: list
    obj_count: int
    first: dt.date
    last: dt.date


def import_source(db: Optional[Db], src: Source, mode: int):
    db = db or Db()
    if mode == MODE_AUTO:
        mode = MODE_REPLACE if db.empty else MODE_APPEND
    if mode == MODE_REPLACE:
        new = Db(list(src.objects), dict(src.groups), dict(src.values))
        rep = ImportReport(mode, src.total, src.total, 0, 0, [], [], len(new.objects),
                           new.first, new.last)
        return new, rep
    new = db.copy()
    keys = {o.upper(): o for o in new.objects}
    new_objects = []
    canon = {}
    for nm in src.objects:
        if nm.upper() in keys:
            canon[nm] = keys[nm.upper()]
            if not new.groups.get(canon[nm]):
                new.groups[canon[nm]] = src.groups.get(nm, "")
        else:
            new.objects.append(nm)
            new.groups[nm] = src.groups.get(nm, "")
            keys[nm.upper()] = nm
            canon[nm] = nm
            new_objects.append(nm)
    added = same = 0
    conflicts = []
    for (d, nm), x in sorted(src.values.items()):
        k = (d, canon[nm])
        if k not in new.values:
            new.values[k] = x
            added += 1
        elif abs(new.values[k] - x) < 1e-7:
            same += 1
        else:
            conflicts.append((canon[nm], d, new.values[k], x))
    rep = ImportReport(mode, src.total, added, same, len(conflicts), conflicts, new_objects,
                       len(new.objects), new.first, new.last)
    return new, rep


# ---------------------------------------------------------------------------
# Dashboard view
# ---------------------------------------------------------------------------

@dataclass
class Cfg:
    green: float = 1.5
    yellow: float = 2.3
    limit: float = 3.0


@dataclass
class ObjView:
    idx: int                 # 1-based base column
    name: str
    group: str
    value_date: Optional[dt.date]
    value: Optional[float]
    prev_date: Optional[dt.date]
    prev_value: Optional[float]
    change: Optional[float]
    zone: int
    stale: bool


@dataclass
class View:
    today: dt.date
    month_start: dt.date
    month_end: dt.date
    is_current: bool
    ref_date: dt.date
    label: str
    objects: list
    sorted: list
    data_date: Optional[dt.date]
    cnt_valued: int
    cnt_green: int
    cnt_yellow: int
    cnt_red: int            # red + over-limit
    cnt_over: int
    cnt_nodata: int
    max_val: Optional[float]
    max_name: str
    avg_val: Optional[float]
    state: int
    obj_count: int
    stale_count: int
    days: list               # dates of the selected month (all days)
    trend: list              # per sorted row: list of 31 (value|None)
    maxday: list             # 31 (value|None)
    last_max_date: Optional[dt.date]
    months: list             # [(start, label)] newest first
    cfg: Cfg


def zone_of(v: Optional[float], cfg: Cfg) -> int:
    if v is None:
        return 0
    if v <= cfg.green:
        return 1
    if v <= cfg.yellow:
        return 2
    if v <= cfg.limit:
        return 3
    return 4


def month_list(db: Db) -> list:
    if db.empty:
        return []
    out = []
    m = month_start(db.last)
    while m >= month_start(db.first):
        out.append((m, month_label(m)))
        m = add_months(m, -1)
    return out


def default_month(db: Db, today: dt.date) -> dt.date:
    """Month the dashboard shows when nothing (valid) is selected."""
    return month_start(db.last) if not db.empty else month_start(today)


def compute(db: Db, today: dt.date, sel: Optional[dt.date] = None, cfg: Cfg = Cfg()) -> View:
    ms = month_start(sel) if sel else default_month(db, today)
    me = month_end(ms)
    is_current = ms <= today <= me
    ref = today if is_current else me
    # current month: today's value, or the previous day's (which on the 1st of a
    # month lies in the previous month); a past month: its last value
    low = min(ms, ref - dt.timedelta(days=1)) if is_current else ms
    objs = []
    names = db.objects[:MAX_OBJ]
    for i, nm in enumerate(names, 1):
        vd = None
        d = ref
        while d >= low:
            if db.get(d, nm) is not None:
                vd = d
                break
            d -= dt.timedelta(days=1)
        val = db.get(vd, nm) if vd else None
        pd = pv = chg = None
        if vd is not None:
            earlier = [d for (d, n) in db.values if n == nm and d < vd]
            if earlier:
                pd = max(earlier)
                pv = db.get(pd, nm)
                chg = round(val - pv, 4)
        objs.append(ObjView(i, nm, db.groups.get(nm, ""), vd, val, pd, pv, chg,
                            zone_of(val, cfg), False))
    data_date = max((o.value_date for o in objs if o.value_date), default=None)
    for o in objs:
        o.stale = o.value is not None and data_date is not None and o.value_date < data_date
    valued = [o for o in objs if o.value is not None]
    nodata = [o for o in objs if o.value is None]
    srt = sorted(valued, key=lambda o: (-o.value, o.idx)) + nodata
    cnt = {z: sum(1 for o in objs if o.zone == z) for z in range(5)}
    if not valued:
        state = 0
    elif cnt[4]:
        state = 4
    elif cnt[3]:
        state = 3
    elif cnt[2]:
        state = 2
    else:
        state = 1
    days = [ms + dt.timedelta(days=i) for i in range((me - ms).days + 1)]
    trend = []
    for o in srt:
        row = []
        for k in range(31):
            if k < len(days) and days[k] <= ref:
                row.append(db.get(days[k], o.name))
            else:
                row.append(None)
        trend.append(row)
    maxday = []
    for k in range(31):
        if k < len(days) and days[k] <= ref:
            vals = [db.get(days[k], nm) for nm in names]
            vals = [v for v in vals if v is not None]
            maxday.append(max(vals) if vals else None)
        else:
            maxday.append(None)
    last_max = max((days[k] for k in range(len(days)) if maxday[k] is not None), default=None)
    return View(
        today=today, month_start=ms, month_end=me, is_current=is_current, ref_date=ref,
        label=month_label(ms), objects=objs, sorted=srt, data_date=data_date,
        cnt_valued=len(valued), cnt_green=cnt[1], cnt_yellow=cnt[2], cnt_red=cnt[3] + cnt[4],
        cnt_over=cnt[4], cnt_nodata=cnt[0], max_val=max((o.value for o in valued), default=None),
        max_name=srt[0].name if valued else "",
        avg_val=(sum(o.value for o in valued) / len(valued)) if valued else None,
        state=state, obj_count=len(names), stale_count=sum(1 for o in objs if o.stale),
        days=days, trend=trend, maxday=maxday, last_max_date=last_max, months=month_list(db),
        cfg=cfg)


# ---------------------------------------------------------------------------
# Expected visible texts on the dashboard (ru-RU number formatting)
# ---------------------------------------------------------------------------

def zone_label(z: int, cfg: Cfg) -> str:
    g, y, lim = (fmt_num(x, 1, 1) for x in (cfg.green, cfg.yellow, cfg.limit))
    return ("—", f"0 – {g} ppm", f"{g} – {y} ppm", f"> {y} ppm", f"> {lim} ppm")[z]


def state_sub(v: View) -> str:
    return ("Загрузите файл с данными", "Превышений не зафиксировано",
            f"В желтой зоне: {v.cnt_yellow} СИКН", f"В красной зоне: {v.cnt_red} СИКН",
            f"Выше норматива: {v.cnt_over} СИКН")[v.state]


def footer(v: View) -> str:
    return "СОСТОЯНИЕ: " + ("НЕТ ДАННЫХ" if v.state == 0 else
                            "ВСЕ В НОРМЕ" if v.state == 1 else "ТРЕБУЕТСЯ ВНИМАНИЕ")


def legend(count: int, valued: int) -> str:
    pct = Decimal(100 * count / max(1, valued)).quantize(Decimal(1), rounding=ROUND_HALF_UP)
    return f"{count} ({pct}%)"


def table_texts(v: View) -> list:
    """24 rows x (№, name, value, status, zone, change, arrow) as displayed."""
    rows = []
    for k in range(TABLE_ROWS):
        if k >= len(v.sorted):
            rows.append(("", "", "", "", "", "", ""))
            continue
        o = v.sorted[k]
        val = "—" if o.value is None else fmt_num(o.value)
        chg = "" if o.change is None else fmt_change(o.change)
        arrow = "" if o.change is None else ("↑" if o.change > 0.00001 else
                                             "↓" if o.change < -0.00001 else "—")
        rows.append((str(k + 1), o.name, val, ROW_STATUS[o.zone], zone_label(o.zone, v.cfg),
                     chg, arrow))
    return rows


def dash_texts(v: View) -> dict:
    """Expected strings of the KPI / header cells, keyed by a logical id."""
    return {
        "state_text": STATE_TEXT[v.state],
        "state_sub": state_sub(v),
        "max": "—" if v.max_val is None else fmt_num(v.max_val),
        "max_name": v.max_name,
        "avg": "—" if v.avg_val is None else fmt_num(v.avg_val),
        "avg_sub": "" if v.cnt_valued == 0 else f"по {v.cnt_valued} СИКН",
        "green": str(v.cnt_green),
        "yellow": str(v.cnt_yellow),
        "red": str(v.cnt_red),
        "total": str(v.obj_count),
        "total_sub": f"нет данных: {v.cnt_nodata}" if v.cnt_nodata > 0 else "",
        "data_date": fmt_date(v.data_date),
        "title": "СОДЕРЖАНИЕ ЛХОС ПО СИКН — " + v.label.upper(),
        "note": ("База пуста — нажмите «ИМПОРТ ДАННЫХ»" if v.obj_count == 0 else
                 "* — значение за предыдущую дату" if v.stale_count else ""),
        "footer": footer(v),
        "leg_g": legend(v.cnt_green, v.cnt_valued),
        "leg_y": legend(v.cnt_yellow, v.cnt_valued),
        "leg_r": legend(v.cnt_red, v.cnt_valued),
        "donut_total": str(v.obj_count),
        "zone_r": zone_label(3, v.cfg),
        "zone_y": zone_label(2, v.cfg),
        "zone_g": zone_label(1, v.cfg),
    }


def report_lines(rep: ImportReport, src: Source) -> list:
    """Lines the import report must contain (substring checks)."""
    lines = [f"Файл: {src.file_name}",
             f"Значений в файле: {rep.total:,}".replace(",", "\xa0"),
             f"Объектов в базе: {rep.obj_count}",
             f"Период данных в базе: {fmt_date(rep.first)} – {fmt_date(rep.last)}"]
    if rep.mode == MODE_REPLACE:
        lines.append("Режим: полная загрузка")
    else:
        lines.append("Режим: добавление новых значений")
        lines.append(f"Уже были в базе: {rep.same:,}".replace(",", "\xa0"))
        if rep.conflicts:
            lines.append(f"Отличаются от базы (оставлены как в базе): {rep.conflicts}")
        if rep.new_objects:
            lines.append(f"Новых объектов: {len(rep.new_objects)} ({', '.join(rep.new_objects)})")
    if src.text_skipped:
        lines.append(f"Пропущено нечисловых значений: {src.text_skipped}")
    return lines


if __name__ == "__main__":  # quick manual check
    import sys
    s = read_source(sys.argv[1])
    db, rep = import_source(None, s, 0)
    print(rep)
    v = compute(db, dt.date.today())
    print(v.label, v.state, v.max_val, v.max_name, v.avg_val, v.data_date)
    for row in table_texts(v):
        print(row)
