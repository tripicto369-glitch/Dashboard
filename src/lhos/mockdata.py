"""Mock source workbooks for the LHOS dashboard ("Содержание ЛХОС по СИКН").

The operator keeps one workbook per year. Every month lives on its own sheet
(named after the month in Russian) laid out exactly like the template the user
supplied, and a "Свод" sheet summarises monthly / yearly averages. Values are
entered once a day.

This module fabricates such workbooks with invented objects and deterministic
pseudo-random values so the dashboard's importer can be exercised end to end:

* ``Мокап_ЛХОС_2026.xlsx`` - the "main" file: data from 2026-01-01 up to
  2026-10-06, an empty pre-prepared "Ноябрь" sheet and the "Свод" sheet.
* ``Мокап_ЛХОС_2026_дополнение.xlsx`` - the "update" file: identical, plus the
  2026-10-07 row, a brand-new object column (W, "Октябрь" only) and one
  silently corrected historical value (to demonstrate append-mode conflicts).

Template layout replicated here (verified against the user's file with openpyxl):

* month sheet: B4:B5 merged "Дата"; row 4 = product-group headers (merged
  across their objects or single cells); row 5 = object names (Arial Cyr 8,
  centred, wrapped, thin borders, row height 22.5); dates in B6.. (builtin
  number format 14, centred, thin borders); values in C6:V.. (General, centred,
  thin borders); column B width 11.71; zoom 85 %; A4 landscape fit-to-page;
  a line chart "Динамика содержания ХОС в товарной нефти по СИКН" under the table.
* "Свод": A2:A3 "Показатель", row 2 groups, row 3 names, row 4 "Среднее
  значение за месяц:", rows 5..16 = Январь..Декабрь with ``AVERAGE`` over each
  month sheet, row 17 "Среднее значение за год" (``AVERAGE`` over all month
  sheets), medium outer borders, number format ``0.00`` and a conditional
  format "> 1" -> yellow fill on B5:U17.

Run ``python3 src/lhos/mockdata.py`` to (re)generate the files in ``dist/``.
"""

from __future__ import annotations

import argparse
import calendar
import datetime as dt
import json
import random
import re
import sys
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Union

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, TwoCellAnchor
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Color, Font, PatternFill, Side
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

# ---------------------------------------------------------------------------
# Calendar / layout constants
# ---------------------------------------------------------------------------

YEAR = 2026
MONTHS_RU = (
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
)
SUMMARY_SHEET = "Свод"

MAIN_LAST_DATE = dt.date(YEAR, 10, 6)     # last day with values in the main file
UPDATE_LAST_DATE = dt.date(YEAR, 10, 7)   # ... and in the update file
LAST_SHEET_MONTH = 11                     # "Ноябрь" exists but is still empty

GROUP_ROW = 4          # product-group header row on month sheets
NAME_ROW = 5           # object-name row on month sheets
FIRST_DATA_ROW = 6     # B6 = 1st day of the month
DATE_COL = "B"
DATE_FORMAT = "mm-dd-yy"   # openpyxl spelling of Excel builtin number format 14

# Svod layout: objects start one column to the left of the month sheets
# (month column C -> Svod column B) because column A holds the row captions.
SVOD_GROUP_ROW = 2
SVOD_NAME_ROW = 3
SVOD_CAPTION_ROW = 4
SVOD_FIRST_MONTH_ROW = 5      # Январь; Декабрь is row 16
SVOD_YEAR_ROW = 17

CHART_TITLE = "Динамика содержания ХОС в товарной нефти по СИКН"

# Zones used by the dashboard (ppm): green <= 1.5 < yellow <= 2.3 < red; limit 3.0.
GREEN_MAX = 1.5
LIMIT = 3.0


@dataclass(frozen=True)
class SourceObject:
    """One measured object (a column of a month sheet)."""

    col: str      # column letter on month sheets, e.g. "C"
    name: str     # header in row 5, e.g. "СИКН №301"
    group: str    # product group shown in row 4


# Product groups in column order: (group header, object names, merged header?).
# Positions and merge shapes mirror the user's template (C4:E4, F4:H4, I4, J4,
# K4:R4, S4:T4, U4:V4).
_GROUP_LAYOUT = (
    ("Нефть тип А", ("СИКН №301", "СИКН №304", "СИКН №307"), True),
    ("Нефть тип Б", ("СИКН №310", "СИКН №312", "СИКН №315"), True),
    ("Нефть тип Б", ("СИКН №318",), False),
    ("Нефть тип Б", ("ПСП Восточный",), False),
    ("Нефть тип А", ("СИКН №320", "СИКН №322", "СИКН №325", "СИКН №327",
                     "СИКН №330", "СИКН №333", "СИКН №336", "СИКН №339"), True),
    ("Нефть тип Б", ("СИКН №341", "СИКН №344"), True),
    ("Узлы подготовки", ("УПСВ-1 Лесная", "УПСВ-2 Озерная"), True),
)


def _build_objects() -> tuple[tuple[SourceObject, ...], tuple[dict, ...]]:
    objects: list[SourceObject] = []
    groups: list[dict] = []
    col = column_index_from_string("C")
    for group, names, merged in _GROUP_LAYOUT:
        first = get_column_letter(col)
        for name in names:
            objects.append(SourceObject(get_column_letter(col), name, group))
            col += 1
        last = get_column_letter(col - 1)
        groups.append({
            "name": group,
            "range": f"{first}{GROUP_ROW}:{last}{GROUP_ROW}" if merged else f"{first}{GROUP_ROW}",
            "merged": merged,
            "objects": list(names),
        })
    return tuple(objects), tuple(groups)


OBJECTS, GROUPS = _build_objects()
OBJECT_BY_NAME = {o.name: o for o in OBJECTS}

# The object that appears only in the update file, on the "Октябрь" sheet.
NEW_OBJECT = SourceObject("W", "СИКН №350", "Нефть тип А")
NEW_OBJECT_MONTH = 10

# ---------------------------------------------------------------------------
# Styles (one constant per distinct cell format of the template)
# ---------------------------------------------------------------------------

_AUTO = Color(indexed=64)                       # "automatic" border colour
_THIN = Side(style="thin", color=_AUTO)
_MEDIUM = Side(style="medium", color=_AUTO)
_NONE = Side()


def _border(left: Side = _THIN, right: Side = _THIN, top: Side = _THIN,
            bottom: Side = _THIN) -> Border:
    return Border(left=left, right=right, top=top, bottom=bottom)


BORDER_THIN = _border()
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_CENTER_WRAP = Alignment(horizontal="center", vertical="center", wrap_text=True)
ALIGN_HCENTER = Alignment(horizontal="center")
ALIGN_WRAP = Alignment(wrap_text=True)
ALIGN_CAPTION = Alignment(horizontal="left", indent=2)

FONT_NAMES = Font(name="Arial Cyr", sz=8, charset=204)        # object names (row 5)
FONT_GROUP_INNER = Font(name="Arial Cyr", sz=10, charset=204)  # covered cells of some merges
# Manual highlighting an operator applies to notable values (template has a few).
FILL_HIGHLIGHT = PatternFill("solid", fgColor="FFFFFF00", bgColor=_AUTO)
# Differential style of the "Свод" conditional format (> 1): yellow background.
FILL_CF = PatternFill(bgColor="FFFFFF00")

# Row-4 header cell formats that differ from the plain "thin box, centred"
# format. They mirror the template cell by cell: the covered cells of the
# K:R and S:T merges carry an Arial Cyr 10 font, and the U:V merge is drawn
# as two half-open boxes. Keys are month-sheet columns (Svod is shifted left).
_GROUP_CELL_OVERRIDES = {
    **{c: "inner" for c in "LMNOPQRT"},
    "U": "open_right",
    "V": "open_left",
}

# Svod: the template draws the right-hand medium border after column T.
SVOD_MEDIUM_RIGHT_COL = "T"
SVOD_LAST_COL = "U"

Value = Union[float, str]   # a stored cell value (str = text typed by the operator)
CellKey = tuple[dt.date, str]  # (date, object name)


def _style(cell, *, font=None, border=None, alignment=None, number_format=None, fill=None):
    """Apply the given style parts to ``cell`` (``None`` keeps the default)."""
    if font is not None:
        cell.font = font
    if border is not None:
        cell.border = border
    if alignment is not None:
        cell.alignment = alignment
    if number_format is not None:
        cell.number_format = number_format
    if fill is not None:
        cell.fill = fill


def _style_group_cell(cell, month_col: str) -> None:
    kind = _GROUP_CELL_OVERRIDES.get(month_col, "box")
    if kind == "inner":
        _style(cell, font=FONT_GROUP_INNER, border=BORDER_THIN, alignment=ALIGN_CENTER)
    elif kind == "open_right":
        _style(cell, border=_border(right=_NONE), alignment=ALIGN_HCENTER)
    elif kind == "open_left":
        _style(cell, border=_border(left=_NONE), alignment=ALIGN_HCENTER)
    else:
        _style(cell, border=BORDER_THIN, alignment=ALIGN_CENTER)


# ---------------------------------------------------------------------------
# Value generation
# ---------------------------------------------------------------------------

def _days(first: dt.date, last: dt.date) -> list[dt.date]:
    return [first + dt.timedelta(days=i) for i in range((last - first).days + 1)]


def _month_days(month: int) -> int:
    return calendar.monthrange(YEAR, month)[1]


def _cell_ref(day: dt.date, obj: SourceObject) -> dict:
    """Location of a (date, object) cell in a source workbook."""
    return {
        "sheet": MONTHS_RU[day.month - 1],
        "cell": f"{obj.col}{FIRST_DATA_ROW + day.day - 1}",
        "date": day.isoformat(),
        "object": obj.name,
    }


def _round_like(rng: random.Random, value: float) -> float:
    """Operators mostly type one decimal; about one value in five has two."""
    return round(value, 2 if rng.random() < 0.2 else 1)


def _decimals(value: float) -> int:
    return 1 if round(value, 1) == value else 2


def _green_series(rng: random.Random, days: list[dt.date]) -> dict[dt.date, float]:
    """Smooth daily series for one object: a baseline plus AR(1) noise, kept green."""
    baseline = rng.uniform(0.1, 1.0)
    drift = 0.0
    series = {}
    for day in days:
        drift = 0.6 * drift + rng.gauss(0.0, 0.07)
        raw = min(max(baseline + drift, 0.05), 1.45)
        series[day] = _round_like(rng, raw)
    return series


@dataclass
class _Dataset:
    """Every value of both files plus the bookkeeping the summary reports."""

    main: dict[CellKey, Value]
    update: dict[CellKey, Value]
    specials: dict
    blanks: list[CellKey]


def _generate(seed: int) -> _Dataset:
    rng = random.Random(seed)
    all_days = _days(dt.date(YEAR, 1, 1), UPDATE_LAST_DATE)
    values: dict[CellKey, Value] = {}
    for obj in OBJECTS:
        for day, value in _green_series(rng, all_days).items():
            values[(day, obj.name)] = value

    def key(month: int, day: int, name: str) -> CellKey:
        return (dt.date(YEAR, month, day), name)

    # -- hand-placed cells the dashboard demo relies on ----------------------
    exceedance = key(3, 17, "СИКН №327")
    text_na = key(2, 10, "СИКН №312")
    numeric_text = key(3, 3, "СИКН №330")
    latest_yellow = key(10, 6, "СИКН №322")
    latest_red = key(10, 6, "СИКН №307")
    latest_blank = key(10, 6, "СИКН №341")
    latest_blank_prev = key(10, 5, "СИКН №341")
    # A month whose last day is missing for one object: the dashboard must fall
    # back to the previous day when showing a past month.
    month_end_blank = key(9, 30, "СИКН №315")
    month_end_prev = key(9, 29, "СИКН №315")
    # Historical value the update file silently corrects (append-mode conflict).
    changed = key(10, 3, "СИКН №301")

    values[exceedance] = 3.2
    values[text_na] = "н/д"
    values[numeric_text] = "0,45"
    values[latest_yellow] = 1.8
    values[latest_red] = 2.5
    values[latest_blank_prev] = 0.6
    del values[latest_blank]
    del values[month_end_blank]
    # Let the two non-green objects of the latest day rise over a few days and
    # ease off on 2026-10-07 (still green), so trends and 24 h deltas look real.
    ramps = {
        "СИКН №322": {4: 1.1, 5: 1.4, 7: 1.4},
        "СИКН №307": {5: 1.3, 7: 1.1},
    }
    for name, by_day in ramps.items():
        for day_no, value in by_day.items():
            values[key(10, day_no, name)] = value

    protected = {exceedance, text_na, numeric_text, latest_yellow, latest_red,
                 latest_blank, latest_blank_prev, month_end_blank, month_end_prev,
                 changed, key(10, 4, "СИКН №322")}
    # The last two days of the main file and the update day stay fully filled
    # (apart from the deliberate blank above), so "latest value" is predictable.
    for day in (dt.date(YEAR, 10, 5), MAIN_LAST_DATE, UPDATE_LAST_DATE):
        protected.update((day, o.name) for o in OBJECTS)

    # -- a few yellow and red days, January..September only --------------------
    spike_days = _days(dt.date(YEAR, 1, 1), dt.date(YEAR, 9, 30))
    candidates = [(d, o.name) for d in spike_days for o in OBJECTS
                  if (d, o.name) not in protected]
    picked = rng.sample(candidates, 10)
    yellow_keys, red_keys = sorted(picked[:6]), sorted(picked[6:])
    for k in yellow_keys:
        values[k] = round(rng.uniform(1.6, 2.2), 1)
    for k in red_keys:
        values[k] = round(rng.uniform(2.4, 2.9), 1)
    protected.update(picked)

    # -- ~2 % missing measurements over the main file's period ----------------
    main_days = _days(dt.date(YEAR, 1, 1), MAIN_LAST_DATE)
    eligible = [(d, o.name) for d in main_days for o in OBJECTS
                if (d, o.name) not in protected]
    total_cells = len(main_days) * len(OBJECTS)
    random_blanks = sorted(rng.sample(eligible, round(0.02 * total_cells) - 2))
    for k in random_blanks:
        del values[k]
    blanks = sorted(random_blanks + [latest_blank, month_end_blank])

    main = {k: v for k, v in values.items() if k[0] <= MAIN_LAST_DATE}

    # -- update file: + 2026-10-07, + new object, + one corrected value -------
    update = dict(values)   # includes 2026-10-07 for every object
    october = _days(dt.date(YEAR, NEW_OBJECT_MONTH, 1), UPDATE_LAST_DATE)
    for day, value in _green_series(rng, october).items():
        update[(day, NEW_OBJECT.name)] = value
    original = main[changed]
    update[changed] = round(original + 0.3, _decimals(original))

    def ref(k: CellKey, value: Value | None) -> dict:
        obj = OBJECT_BY_NAME.get(k[1], NEW_OBJECT)
        return {**_cell_ref(k[0], obj), "value": value}

    specials = {
        "exceedance": ref(exceedance, 3.2),
        "text_value": ref(text_na, "н/д"),
        "numeric_text": {**ref(numeric_text, "0,45"), "parsed": 0.45},
        "yellow_days": [ref(k, values[k]) for k in yellow_keys],
        "red_days": [ref(k, values[k]) for k in red_keys],
        "latest_day": {
            "date": MAIN_LAST_DATE.isoformat(),
            "yellow": ref(latest_yellow, 1.8),
            "red": ref(latest_red, 2.5),
            "blank": ref(latest_blank, None),
            "blank_previous": ref(latest_blank_prev, 0.6),
        },
        "month_end_blank": {
            "blank": ref(month_end_blank, None),
            "previous": ref(month_end_prev, values[month_end_prev]),
        },
        "update_changed": {**ref(changed, update[changed]), "original": original},
        "update_new_object": {
            "col": NEW_OBJECT.col, "name": NEW_OBJECT.name, "group": NEW_OBJECT.group,
            "sheet": MONTHS_RU[NEW_OBJECT_MONTH - 1],
            "header_cells": [f"{NEW_OBJECT.col}{GROUP_ROW}", f"{NEW_OBJECT.col}{NAME_ROW}"],
            "first_date": october[0].isoformat(), "last_date": october[-1].isoformat(),
        },
    }
    return _Dataset(main=main, update=update, specials=specials, blanks=blanks)


# ---------------------------------------------------------------------------
# Workbook writing
# ---------------------------------------------------------------------------

def _page_setup(ws, *, landscape: bool, fit: bool) -> None:
    ws.page_margins = PageMargins(left=0.7, right=0.7, top=0.75, bottom=0.75,
                                  header=0.3, footer=0.3)
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    if fit:
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        ws.page_setup.scale = 52


def _add_trend_chart(ws, objects: list[SourceObject], last_row: int) -> None:
    """Line chart below the table, like the one in the template's month sheet."""
    chart = LineChart()
    chart.title = CHART_TITLE
    chart.style = 2
    chart.legend.position = "b"
    chart.display_blanks = "gap"
    chart.y_axis.crossAx = 500
    chart.x_axis = DateAxis(crossAx=100)
    chart.x_axis.axPos = "b"
    chart.x_axis.majorTickMark = "out"
    chart.x_axis.number_format = "m/d/yyyy"
    chart.x_axis.baseTimeUnit = "days"
    chart.x_axis.delete = False      # openpyxl >= 3.1 hides axes unless told otherwise
    chart.y_axis.delete = False
    first_col = column_index_from_string(objects[0].col)
    last_col = column_index_from_string(objects[-1].col)
    chart.add_data(Reference(ws, min_col=first_col, max_col=last_col,
                             min_row=NAME_ROW, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=column_index_from_string(DATE_COL),
                                   min_row=FIRST_DATA_ROW, max_row=last_row))
    for series in chart.series:
        series.smooth = False
        series.marker.symbol = "none"
        series.graphicalProperties.line.width = 28575   # 2.25 pt, as in the template
        series.graphicalProperties.line.cap = "rnd"
    # Same placement as the template: A38 (+offset) .. X63 (+offset).
    anchor = TwoCellAnchor()
    anchor._from = AnchorMarker(col=0, colOff=593912, row=37, rowOff=179294)
    anchor.to = AnchorMarker(col=23, colOff=301758, row=62, rowOff=142555)
    ws.add_chart(chart, anchor)


def _write_month_sheet(ws, month: int, objects: list[SourceObject],
                       values: dict[CellKey, Value]) -> None:
    ws.column_dimensions[DATE_COL].width = 11.7109375
    ws.sheet_view.zoomScale = 85
    ws.sheet_view.zoomScaleNormal = 85
    _page_setup(ws, landscape=True, fit=True)

    # Header: "Дата" spans rows 4-5, product groups in row 4, objects in row 5.
    ws.merge_cells(f"{DATE_COL}{GROUP_ROW}:{DATE_COL}{NAME_ROW}")
    ws[f"{DATE_COL}{GROUP_ROW}"] = "Дата"
    for row in (GROUP_ROW, NAME_ROW):
        _style(ws[f"{DATE_COL}{row}"], border=BORDER_THIN, alignment=ALIGN_CENTER)

    for group in GROUPS:
        if group["merged"]:
            ws.merge_cells(group["range"])
        ws[group["range"].split(":")[0]] = group["name"]
    for obj in objects:
        if obj not in OBJECTS:          # object added later: single-cell group header
            ws[f"{obj.col}{GROUP_ROW}"] = obj.group
        _style_group_cell(ws[f"{obj.col}{GROUP_ROW}"], obj.col)
        name_cell = ws[f"{obj.col}{NAME_ROW}"]
        name_cell.value = obj.name
        _style(name_cell, font=FONT_NAMES, border=BORDER_THIN, alignment=ALIGN_CENTER_WRAP)
    ws.row_dimensions[NAME_ROW].height = 22.5

    # One row per calendar day; empty cells keep their borders like the template.
    days = _month_days(month)
    last_col = column_index_from_string(objects[-1].col)
    for day_no in range(1, days + 1):
        day = dt.date(YEAR, month, day_no)
        row = FIRST_DATA_ROW + day_no - 1
        date_cell = ws[f"{DATE_COL}{row}"]
        # Format first: otherwise openpyxl registers its own "yyyy-mm-dd" format.
        _style(date_cell, number_format=DATE_FORMAT, border=BORDER_THIN, alignment=ALIGN_CENTER)
        date_cell.value = day
        for obj in objects:
            cell = ws[f"{obj.col}{row}"]
            value = values.get((day, obj.name))
            cell.value = value
            highlight = isinstance(value, float) and value > GREEN_MAX
            _style(cell, border=BORDER_THIN, alignment=ALIGN_CENTER,
                   fill=FILL_HIGHLIGHT if highlight else None)
        # The template centres (without borders) the two columns right of the table.
        for col in (last_col + 1, last_col + 2):
            ws.cell(row=row, column=col).alignment = ALIGN_CENTER

    _add_trend_chart(ws, objects, FIRST_DATA_ROW + days - 1)


def _svod_col(obj: SourceObject) -> str:
    return get_column_letter(column_index_from_string(obj.col) - 1)


def _write_svod_sheet(ws, sheet_months: list[int]) -> None:
    ws.column_dimensions["A"].width = 19.140625
    ws.column_dimensions["B"].width = 9.140625
    for col in range(3, column_index_from_string(SVOD_LAST_COL) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 9.5703125
    for row, height in ((1, 15.75), (3, 23.25), (4, 30), (16, 15.75), (17, 30.75)):
        ws.row_dimensions[row].height = height
    _page_setup(ws, landscape=False, fit=False)

    # Header block.
    ws.merge_cells(f"A{SVOD_GROUP_ROW}:A{SVOD_NAME_ROW}")
    ws[f"A{SVOD_GROUP_ROW}"] = "Показатель"
    _style(ws[f"A{SVOD_GROUP_ROW}"], border=_border(left=_MEDIUM, top=_MEDIUM, bottom=_NONE),
           alignment=ALIGN_CENTER)
    _style(ws[f"A{SVOD_NAME_ROW}"], border=_border(left=_MEDIUM, top=_NONE, bottom=_MEDIUM),
           alignment=ALIGN_CENTER)
    for group in GROUPS:
        cols = [_svod_col(OBJECT_BY_NAME[n]) for n in group["objects"]]
        if group["merged"]:
            ws.merge_cells(f"{cols[0]}{SVOD_GROUP_ROW}:{cols[-1]}{SVOD_GROUP_ROW}")
        ws[f"{cols[0]}{SVOD_GROUP_ROW}"] = group["name"]
    for obj in OBJECTS:
        col = _svod_col(obj)
        _style_group_cell(ws[f"{col}{SVOD_GROUP_ROW}"], obj.col)
        ws[f"{col}{SVOD_NAME_ROW}"] = obj.name
        _style(ws[f"{col}{SVOD_NAME_ROW}"], font=FONT_NAMES, border=BORDER_THIN,
               alignment=ALIGN_CENTER_WRAP)

    # Row captions (column A) and the grid; medium outer frame, as in the template.
    ws[f"A{SVOD_CAPTION_ROW}"] = "Среднее значение за месяц:"
    ws[f"A{SVOD_YEAR_ROW}"] = "Среднее значение за год"
    last_month_row = SVOD_FIRST_MONTH_ROW + 11
    for row in range(SVOD_CAPTION_ROW, SVOD_YEAR_ROW + 1):
        top = _MEDIUM if row in (SVOD_CAPTION_ROW, SVOD_YEAR_ROW) else _THIN
        bottom = _MEDIUM if row in (last_month_row, SVOD_YEAR_ROW) else _THIN
        caption = ws[f"A{row}"]
        if row in (SVOD_CAPTION_ROW, SVOD_YEAR_ROW):
            _style(caption, border=_border(left=_MEDIUM, top=top, bottom=bottom),
                   alignment=ALIGN_WRAP)
        else:
            caption.value = MONTHS_RU[row - SVOD_FIRST_MONTH_ROW]
            _style(caption, border=_border(left=_MEDIUM, top=top, bottom=bottom),
                   alignment=ALIGN_CAPTION)
        for obj in OBJECTS:
            col = _svod_col(obj)
            right = _MEDIUM if col == SVOD_MEDIUM_RIGHT_COL else _THIN
            cell = ws[f"{col}{row}"]
            _style(cell, border=_border(right=right, top=top, bottom=bottom))
            if row != SVOD_CAPTION_ROW:
                _style(cell, number_format="0.00", alignment=ALIGN_CENTER)

    # Formulas: one AVERAGE per existing month sheet, plus the yearly average.
    for obj in OBJECTS:
        col = _svod_col(obj)
        year_parts = []
        for month in sheet_months:
            sheet = MONTHS_RU[month - 1]
            last = FIRST_DATA_ROW + _month_days(month) - 1
            ws[f"{col}{SVOD_FIRST_MONTH_ROW + month - 1}"] = (
                f"=AVERAGE({sheet}!{obj.col}${FIRST_DATA_ROW}:{obj.col}${last})")
            year_parts.append(f"{sheet}!{obj.col}{FIRST_DATA_ROW}:{obj.col}{last}")
        ws[f"{col}{SVOD_YEAR_ROW}"] = f"=AVERAGE({','.join(year_parts)})"

    ws.conditional_formatting.add(
        f"B{SVOD_FIRST_MONTH_ROW}:{SVOD_LAST_COL}{SVOD_YEAR_ROW}",
        CellIsRule(operator="greaterThan", formula=["1"], fill=FILL_CF))


def _defined_name(obj: SourceObject) -> str:
    """Workbook name for an object's January column ("УПСВ-1 Лесная" -> "УПСВ_1_Лесная")."""
    return re.sub(r"[\s\-]", "_", obj.name)


def _write_workbook(path: Path, values: dict[CellKey, Value],
                    extra_objects: dict[int, list[SourceObject]],
                    saved_at: dt.datetime) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    sheet_months = list(range(1, LAST_SHEET_MONTH + 1))
    for month in sheet_months:
        ws = wb.create_sheet(MONTHS_RU[month - 1])
        _write_month_sheet(ws, month, list(OBJECTS) + extra_objects.get(month, []), values)
    _write_svod_sheet(wb.create_sheet(SUMMARY_SHEET), sheet_months)

    # The template defines one name per object pointing at its January column.
    for obj in OBJECTS:
        name = _defined_name(obj)
        wb.defined_names[name] = DefinedName(
            name, attr_text=f"{MONTHS_RU[0]}!${obj.col}${FIRST_DATA_ROW}:${obj.col}$"
                            f"{FIRST_DATA_ROW + _month_days(1) - 1}")

    # Open on the month the operator is currently filling in.
    active = UPDATE_LAST_DATE.month - 1
    for index, ws in enumerate(wb.worksheets):
        ws.sheet_view.tabSelected = index == active
    wb.active = active

    wb.properties.creator = "Генератор мокапа ЛХОС"
    wb.properties.lastModifiedBy = "Генератор мокапа ЛХОС"
    wb.properties.title = f"Содержание ЛХОС по СИКН, {YEAR} (тестовые данные)"
    wb.properties.created = dt.datetime(YEAR, 1, 1, 8, 0)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    _finalize_package(path, saved_at)


# openpyxl 3.1 stores every text cell as an inline string; Excel (and the
# user's template) keep them in xl/sharedStrings.xml. Rewriting the package
# keeps the mock structurally identical to an Excel-saved file and makes the
# output byte-for-byte reproducible (fixed zip timestamps and dcterms:modified).
_INLINE_STRING_RE = re.compile(
    r'<c ([^>]*?)t="inlineStr"([^>]*)><is>(<t(?:\s[^>]*)?>.*?</t>)</is></c>', re.S)
_MODIFIED_RE = re.compile(r'(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)')
_SST_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
_SST_REL_TYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"


def _finalize_package(path: Path, saved_at: dt.datetime) -> None:
    with zipfile.ZipFile(path) as src:
        parts = {info.filename: src.read(info) for info in src.infolist()}

    strings: dict[str, int] = {}     # serialized <t> element -> index in the table
    total = 0

    def to_shared(match: re.Match) -> str:
        nonlocal total
        total += 1
        index = strings.setdefault(match.group(3), len(strings))
        return f'<c {match.group(1)}t="s"{match.group(2)}><v>{index}</v></c>'

    # Only convert when openpyxl wrote no string table of its own: otherwise the
    # new indices would point into a table that is never written. Inline
    # strings are valid OOXML, so leaving them alone is the safe fallback.
    if "xl/sharedStrings.xml" not in parts:
        for name in sorted(parts):
            if name.startswith("xl/worksheets/sheet") and name.endswith(".xml"):
                xml = parts[name].decode("utf-8")
                parts[name] = _INLINE_STRING_RE.sub(to_shared, xml).encode("utf-8")
    if strings:
        items = "".join(f"<si>{t}</si>" for t in strings)
        parts["xl/sharedStrings.xml"] = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
            '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
            f'count="{total}" uniqueCount="{len(strings)}">{items}</sst>').encode("utf-8")
        types = parts["[Content_Types].xml"].decode("utf-8")
        parts["[Content_Types].xml"] = types.replace(
            "</Types>", f'<Override PartName="/xl/sharedStrings.xml" '
                        f'ContentType="{_SST_CONTENT_TYPE}"/></Types>').encode("utf-8")
        rels = parts["xl/_rels/workbook.xml.rels"].decode("utf-8")
        parts["xl/_rels/workbook.xml.rels"] = rels.replace(
            "</Relationships>", f'<Relationship Type="{_SST_REL_TYPE}" '
                                f'Target="sharedStrings.xml" Id="rIdSst"/></Relationships>'
        ).encode("utf-8")

    stamp = saved_at.strftime("%Y-%m-%dT%H:%M:%SZ")
    core = parts["docProps/core.xml"].decode("utf-8")
    parts["docProps/core.xml"] = _MODIFIED_RE.sub(rf"\g<1>{stamp}\g<2>", core).encode("utf-8")

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in parts.items():
            info = zipfile.ZipInfo(name, date_time=saved_at.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, data)


# ---------------------------------------------------------------------------
# Summary helpers
# ---------------------------------------------------------------------------

def _is_number(value: Value | None) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _latest(values: dict[CellKey, Value], names: list[str],
            last_date: dt.date, first_date: dt.date | None = None) -> dict:
    """Last numeric value per object within [first_date, last_date]."""
    result = {}
    for name in names:
        dates = [d for (d, n), v in values.items()
                 if n == name and _is_number(v) and d <= last_date
                 and (first_date is None or d >= first_date)]
        if dates:
            day = max(dates)
            result[name] = {"date": day.isoformat(), "value": values[(day, name)]}
    return result


def _file_summary(path: Path, values: dict[CellKey, Value], last_date: dt.date,
                  names: list[str]) -> dict:
    month_last = {}
    for month in range(1, last_date.month + 1):
        first = dt.date(YEAR, month, 1)
        last = min(dt.date(YEAR, month, _month_days(month)), last_date)
        month_last[MONTHS_RU[month - 1]] = _latest(values, names, last, first)
    return {
        "path": str(path),
        "sheets": list(MONTHS_RU[:LAST_SHEET_MONTH]) + [SUMMARY_SHEET],
        "date_first": min(d for d, _ in values).isoformat(),
        "date_last": max(d for d, _ in values).isoformat(),
        "numeric_values": sum(1 for v in values.values() if _is_number(v)),
        "text_values": sum(1 for v in values.values() if isinstance(v, str)),
        "latest": _latest(values, names, last_date),
        "month_last": month_last,
    }


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def make_mock(path_main: str, path_update: str, seed: int = 42) -> dict:
    """Write the main and the update mock workbooks and return a JSON-able summary.

    The summary lists the objects and groups, per-file date ranges and value
    counts, the latest / month-end values per object and every hand-placed
    special cell (sheet, cell address, date, object, value).
    """
    data = _generate(seed)
    main_path, update_path = Path(path_main), Path(path_update)
    # "Saved" shortly after the operator's daily entry on the file's last day.
    _write_workbook(main_path, data.main, {}, dt.datetime(YEAR, 10, 6, 9, 15))
    _write_workbook(update_path, data.update, {NEW_OBJECT_MONTH: [NEW_OBJECT]},
                    dt.datetime(YEAR, 10, 7, 9, 10))

    names = [o.name for o in OBJECTS]
    main = _file_summary(main_path, data.main, MAIN_LAST_DATE, names)
    main["blank_cells"] = [_cell_ref(d, OBJECT_BY_NAME[n]) for d, n in data.blanks]
    update = _file_summary(update_path, data.update, UPDATE_LAST_DATE,
                           names + [NEW_OBJECT.name])
    return {
        "seed": seed,
        "year": YEAR,
        "objects": [{"col": o.col, "name": o.name, "group": o.group} for o in OBJECTS],
        "groups": [dict(g) for g in GROUPS],
        "month_sheets": list(MONTHS_RU[:LAST_SHEET_MONTH]),
        "empty_month_sheets": [MONTHS_RU[LAST_SHEET_MONTH - 1]],
        "summary_sheet": SUMMARY_SHEET,
        "zones": {"green_max": GREEN_MAX, "yellow_max": 2.3, "limit": LIMIT},
        "main": main,
        "update": update,
        "special": data.specials,
    }


def main(argv: list[str] | None = None) -> int:
    repo_root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description="Generate mock LHOS source workbooks.")
    parser.add_argument("--out-dir", type=Path, default=repo_root / "dist")
    parser.add_argument("--seed", type=int, default=42)
    parser.add_argument("--summary-json", type=Path,
                        help="also write the summary dict to this JSON file")
    args = parser.parse_args(argv)

    summary = make_mock(str(args.out_dir / "Мокап_ЛХОС_2026.xlsx"),
                        str(args.out_dir / "Мокап_ЛХОС_2026_дополнение.xlsx"),
                        seed=args.seed)
    for key in ("main", "update"):
        info = summary[key]
        print(f"{info['path']}: {info['date_first']} .. {info['date_last']}, "
              f"{info['numeric_values']} numeric + {info['text_values']} text values")
    if args.summary_json:
        args.summary_json.write_text(json.dumps(summary, ensure_ascii=False, indent=2),
                                     encoding="utf-8")
        print(f"summary: {args.summary_json}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
